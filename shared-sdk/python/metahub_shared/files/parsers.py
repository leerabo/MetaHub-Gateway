"""PDF、DOCX 和 UTF-8 TXT 的有界文本抽取器。

第三方解析器只接收已通过文件大小、扩展名、MIME 和魔数校验的 bytes。本模块逐步累计
正文并在超限时立即失败，任何第三方异常都折叠为固定文件错误，不暴露输入或底层消息。
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field
from pathlib import PurePosixPath
from typing import Final
from xml.etree import ElementTree

from .contracts import (
    EXTRACTOR_VERSION_BY_KIND,
    MAX_EXTRACTED_CODE_POINTS,
    MAX_EXTRACTED_TEXT_BYTES,
    ExtractedText,
    FileContractError,
    build_extracted_text,
)

# DOCX/PDF 容器上限限制解压放大、对象数量和单次解析成本；这些值不能由请求正文覆盖。
MAX_DOCX_ENTRIES: Final[int] = 2_048
MAX_DOCX_UNCOMPRESSED_BYTES: Final[int] = 64 * 1024 * 1024
MAX_DOCX_ENTRY_BYTES: Final[int] = 32 * 1024 * 1024
MAX_DOCX_COMPRESSION_RATIO: Final[int] = 200
MAX_PDF_PAGES: Final[int] = 2_000

# OOXML 核心对象和危险嵌入使用规范化小写路径比较，避免大小写变体绕过检查。
_DOCX_REQUIRED_ENTRIES: Final[frozenset[str]] = frozenset(
    {"[content_types].xml", "word/document.xml"}
)
_DOCX_ARCHIVE_SUFFIXES: Final[frozenset[str]] = frozenset(
    {".7z", ".bz2", ".gz", ".rar", ".tar", ".tgz", ".xz", ".zip"}
)
_RELATIONSHIP_TAG: Final[str] = "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"


@dataclass(slots=True)
class _TextAccumulator:
    """逐段累计正文并在 join 前执行 code-point/UTF-8 双重上限。

    该对象只存在于解析调用栈，``parts`` 关闭 repr；每次追加都更新准确计数，防止先构造
    无界页面列表再由最终 DTO 拒绝。
    """

    parts: list[str] = field(default_factory=list, repr=False)
    code_points: int = 0
    utf8_bytes: int = 0

    def append(self, value: object) -> None:
        """追加一个非空文本段，并在越界前抛出固定错误。

        参数：解析器产生的单段文本。返回值：无。空段被忽略；非字符串和超限值失败关闭。
        """

        if type(value) is not str:
            # 第三方解析器必须返回明确文本，禁止把对象 repr 当正文继续累计。
            raise FileContractError("file.parse_failed")
        if not value:
            return
        separator_size = 1 if self.parts else 0
        encoded_size = len(value.encode("utf-8"))
        next_code_points = self.code_points + separator_size + len(value)
        next_utf8_bytes = self.utf8_bytes + separator_size + encoded_size
        if (
            next_code_points > MAX_EXTRACTED_CODE_POINTS
            or next_utf8_bytes > MAX_EXTRACTED_TEXT_BYTES
        ):
            # 逐段失败避免 join 产生第二份超限文本，错误对象不包含当前段内容。
            raise FileContractError("file.extract_limit_exceeded")
        self.parts.append(value)
        self.code_points = next_code_points
        self.utf8_bytes = next_utf8_bytes

    def finish(self, *, extractor_version: str) -> ExtractedText:
        """合并有界段并生成关闭 repr 的抽取 DTO。"""

        return build_extracted_text("\n".join(self.parts), extractor_version=extractor_version)


def extract_file_text(content: object, kind: object) -> ExtractedText:
    """按可信文件类型分派文本抽取。

    参数：已受上传上限保护的 bytes 和三方格式校验得到的 kind。返回：有界抽取正文。
    关键约束：未知类型不能退化为 TXT 猜测，第三方异常必须由具体解析器固定映射。
    """

    if type(content) is not bytes or kind not in EXTRACTOR_VERSION_BY_KIND:
        # 分派只接受格式门禁产生的 bytes/kind，未知类型不能降级猜测解析器。
        raise FileContractError("file.parse_failed")
    if kind == "txt":
        return _extract_txt(content)
    if kind == "docx":
        return _extract_docx(content)
    if kind == "pdf":
        return _extract_pdf(content)
    # Literal 理论上已穷尽；运行时仍失败关闭，避免未来新增类型被静默当成另一格式。
    raise FileContractError("file.unsupported_format")


def _extract_txt(content: bytes) -> ExtractedText:
    """严格解码 UTF-8/UTF-8 BOM TXT，不执行任何标记语言解释。"""

    try:
        text = content.decode("utf-8-sig", errors="strict")
    # 严格 UTF-8 解码失败不能回显字节位置或原始内容，只返回固定结构错误。
    except UnicodeDecodeError:
        # 解码异常位置和原始 bytes 不进入公开异常。
        raise FileContractError("file.structure_invalid") from None
    if "\x00" in text:
        # 解码后的 NUL 仍代表二进制或含糊文本，不能进入 DLP 和模型处理链路。
        raise FileContractError("file.structure_invalid")
    return build_extracted_text(text, extractor_version=EXTRACTOR_VERSION_BY_KIND["txt"])


def _extract_docx(content: bytes) -> ExtractedText:
    """校验 OOXML ZIP 结构后提取段落、表格、页眉和页脚文本。

    宏、OLE、外部模板/媒体和嵌套包会在交给 python-docx 前拒绝；普通 hyperlink 仅保留
    已存在的显示文本且绝不访问目标 URL。
    """

    _validate_docx_package(content)
    try:
        from docx import Document

        document = Document(io.BytesIO(content))
        accumulator = _TextAccumulator()
        for paragraph in document.paragraphs:
            accumulator.append(paragraph.text)
        for table in document.tables:
            _append_table_text(accumulator, table)
        for section in document.sections:
            # 页眉/页脚属于用户可见 Word 文本；同一链接 section 的重复内容保持可解释顺序。
            for paragraph in section.header.paragraphs:
                accumulator.append(paragraph.text)
            for paragraph in section.footer.paragraphs:
                accumulator.append(paragraph.text)
        return accumulator.finish(extractor_version=EXTRACTOR_VERSION_BY_KIND["docx"])
    # 已识别的资源或结构错误保持原码，不能被 python-docx 的兜底异常覆盖。
    except FileContractError:
        raise
    except Exception:
        # python-docx 异常可能携带 ZIP entry 或 XML 文本，只暴露固定解析错误。
        raise FileContractError("file.parse_failed") from None


def _validate_docx_package(content: bytes) -> None:
    """检查 DOCX entry、解压放大、危险嵌入和外部关系。

    参数：完整但有界的 DOCX bytes。返回值：无。关键约束：只读检查，不把 entry 名称或
    关系 URL 写入异常；所有 XML 都来自受单 entry/总解压上限保护的包。
    """

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if not entries or len(entries) > MAX_DOCX_ENTRIES:
                # 空包或 entry 数超限都无法形成可信且有界的 Word 文档。
                raise FileContractError("file.extract_limit_exceeded")
            normalized_entries: dict[str, str] = {}
            total_uncompressed = 0
            for entry in entries:
                normalized = _normalize_docx_entry(entry.filename)
                if normalized in normalized_entries:
                    # ZIP 名称大小写可不同但 OOXML 语义会冲突，拒绝双重解释而不是任选其一。
                    raise FileContractError("file.structure_invalid")
                normalized_entries[normalized] = entry.filename
                if entry.file_size < 0 or entry.file_size > MAX_DOCX_ENTRY_BYTES:
                    # 单 entry 上限在读取 XML 前执行，防止压缩包放大占用内存。
                    raise FileContractError("file.extract_limit_exceeded")
                total_uncompressed += entry.file_size
                if total_uncompressed > MAX_DOCX_UNCOMPRESSED_BYTES:
                    # 总解压体积独立受限，多个小 entry 不能绕过单项门禁。
                    raise FileContractError("file.extract_limit_exceeded")
                if entry.file_size > 0 and (
                    entry.compress_size <= 0
                    or entry.file_size > entry.compress_size * MAX_DOCX_COMPRESSION_RATIO
                ):
                    # 异常压缩比按潜在 ZIP bomb 拒绝，不把 entry 名称写入错误。
                    raise FileContractError("file.extract_limit_exceeded")
                _reject_dangerous_docx_entry(normalized)
            if not _DOCX_REQUIRED_ENTRIES <= set(normalized_entries):
                # 缺少内容类型或正文 XML 的 ZIP 不能冒充 DOCX 成功解析。
                raise FileContractError("file.structure_invalid")
            for normalized, original in normalized_entries.items():
                if normalized.endswith(".rels"):
                    # 回读必须使用 ZIP 中的原始大小写名称，规范名只用于安全比较和去重。
                    _validate_relationships(archive.read(original))
    # 本模块固定错误保持原语义，其余 ZIP/I/O 异常统一折叠为结构不可用。
    except FileContractError:
        raise
    # ZIP 异常可能携带内部 entry 名，必须在离开解析边界前清除底层消息。
    except (KeyError, OSError, RuntimeError, zipfile.BadZipFile, zipfile.LargeZipFile):
        raise FileContractError("file.structure_invalid") from None


def _normalize_docx_entry(value: object) -> str:
    """把 ZIP entry 转为安全小写 POSIX 相对路径，拒绝路径穿越与编码歧义。"""

    if type(value) is not str or not value or "\\" in value or "\x00" in value:
        # ZIP entry 必须是非空 POSIX 文本，反斜杠和 NUL 会产生平台相关解释。
        raise FileContractError("file.structure_invalid")
    path = PurePosixPath(value)
    if path.is_absolute() or any(part in {"", ".", ".."} for part in path.parts):
        # entry 只能在包内寻址，禁止 zip-slip 和含糊的当前目录片段。
        raise FileContractError("file.structure_invalid")
    if path.parts and ":" in path.parts[0]:
        # Windows 盘符或 alternate stream 语法不能出现在 OOXML 包内路径。
        raise FileContractError("file.structure_invalid")
    return path.as_posix().lower()


def _reject_dangerous_docx_entry(name: str) -> None:
    """拒绝宏、OLE、嵌套压缩包和可执行嵌入 entry。"""

    suffix = PurePosixPath(name).suffix.lower()
    if (
        name.endswith("vbaproject.bin")
        or "/embeddings/" in f"/{name}"
        or "oleobject" in name
        or suffix in _DOCX_ARCHIVE_SUFFIXES
    ):
        # 文档正文解析不需要这些对象；保留会扩大 PRD-073 前的攻击面。
        raise FileContractError("file.structure_invalid")


def _validate_relationships(content: bytes) -> None:
    """允许普通外部 hyperlink，拒绝会触发模板、媒体或数据访问的外部关系。"""

    try:
        root = ElementTree.fromstring(content)
    # 关系 XML 损坏时不继续猜测，避免解析器对外部目标产生不同解释。
    except ElementTree.ParseError:
        raise FileContractError("file.structure_invalid") from None
    for relationship in root.iter(_RELATIONSHIP_TAG):
        if relationship.attrib.get("TargetMode", "").lower() != "external":
            continue
        relation_type = relationship.attrib.get("Type", "").lower()
        if not relation_type.endswith("/hyperlink"):
            # 目标值可能含内网地址或凭据，拒绝时不得拼入异常。
            raise FileContractError("file.structure_invalid")


def _append_table_text(accumulator: _TextAccumulator, table: object) -> None:
    """按行列顺序提取 python-docx 表格单元格文本。

    参数：当前 accumulator 与 python-docx Table。返回值：无。只读取公开 rows/cells/text
    投影；结构损坏交给外层固定解析错误处理。
    """

    rows = table.rows
    for row in rows:
        for cell in row.cells:
            accumulator.append(cell.text)


def _extract_pdf(content: bytes) -> ExtractedText:
    """使用 pypdf 逐页抽取文本并保持页数、解压和正文上限。

    加密、损坏或无文本 PDF 不会降级到 OCR；JavaScript、附件和外部 URI 也不会被执行。
    """

    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content), strict=True)
        if reader.is_encrypted:
            # 本 PRD 没有口令输入合同，加密 PDF 不能静默尝试空密码或降级解析。
            raise FileContractError("file.structure_invalid")
        page_count = len(reader.pages)
        if page_count > MAX_PDF_PAGES:
            # 页数上限在逐页抽取前执行，防止恶意页树放大同步处理时间。
            raise FileContractError("file.extract_limit_exceeded")
        accumulator = _TextAccumulator()
        for page in reader.pages:
            # pypdf 受运行依赖自身解压上限保护；本层在每页返回后立即执行更小的正文上限。
            accumulator.append(page.extract_text() or "")
        return accumulator.finish(extractor_version=EXTRACTOR_VERSION_BY_KIND["pdf"])
    # 已识别的加密、页数或正文上限错误保持原码，便于调用方采取正确动作。
    except FileContractError:
        raise
    except Exception:
        # pypdf 可能抛出包含对象号或原始 token 的异常，统一折叠为固定错误。
        raise FileContractError("file.parse_failed") from None
