"""PRD-072 文件格式、解析、DLP、密文和本地 Blob Store 专项测试。"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from metahub_shared.files import (
    DEFAULT_FILE_ENCRYPTION_SECRET,
    FILE_DLP_WINDOW_CODE_POINTS,
    EncryptedLocalBlobStore,
    FileBlobOwner,
    FileContractError,
    FileDlpRule,
    decrypt_file_blob,
    encrypt_file_blob,
    estimate_file_tokens,
    extract_file_text,
    file_content_fingerprint,
    file_object_key,
    normalize_filename,
    scan_file_text,
    validate_file_secret,
    validate_upload_format,
)
from metahub_shared.runtime.config import ConfigError, load_app_config


def _pdf_with_text(text: str) -> bytes:
    """生成不依赖额外测试库的单页 Helvetica 文本 PDF。"""

    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        (
            b"<< /Length "
            + str(len(stream)).encode("ascii")
            + b" >>\nstream\n"
            + stream
            + b"\nendstream"
        ),
    ]
    payload = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for index, value in enumerate(objects, start=1):
        offsets.append(len(payload))
        payload.extend(f"{index} 0 obj\n".encode("ascii"))
        payload.extend(value)
        payload.extend(b"\nendobj\n")
    xref = len(payload)
    payload.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    payload.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref}\n%%EOF\n"
        ).encode("ascii")
    )
    return bytes(payload)


def _docx_with_text(*, paragraph: str = "合同正文", cell: str = "表格字段") -> bytes:
    """生成包含段落和表格的最小 DOCX 测试文件。"""

    from docx import Document

    document = Document()
    document.add_paragraph(paragraph)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = cell
    target = io.BytesIO()
    document.save(target)
    return target.getvalue()


def _rule(
    rule_id: str,
    *,
    pattern: str,
    action: str,
    order: int,
    rule_type: str = "keyword",
) -> FileDlpRule:
    """构造顺序和摘要固定的文件 DLP 测试规则。"""

    return FileDlpRule(
        dlp_rule_id=rule_id,
        rule_type=rule_type,  # type: ignore[arg-type]
        pattern=pattern,
        pattern_hash=f"sha256:{order:064x}",
        action=action,  # type: ignore[arg-type]
        priority=order + 1,
        updated_at=f"2026-08-05T00:00:0{order}Z",
        order=order,
    )


def test_filename_and_format_require_exact_extension_mime_and_magic() -> None:
    """主格式门禁拒绝路径、危险复合后缀和 MIME/魔数错配。"""

    assert normalize_filename("报告.PDF") == ("报告.PDF", "pdf")
    result = validate_upload_format(
        filename="报告.PDF",
        media_type="application/pdf",
        content=_pdf_with_text("hello"),
    )
    assert (result.kind, result.extension, result.media_type) == (
        "pdf",
        "pdf",
        "application/pdf",
    )
    for name in ("../a.pdf", "a.exe.pdf", "a\u202epdf.txt", "a.doc"):
        with pytest.raises(FileContractError):
            normalize_filename(name)
    with pytest.raises(FileContractError, match="file.unsupported_format"):
        validate_upload_format(
            filename="a.pdf", media_type="text/plain", content=_pdf_with_text("x")
        )
    with pytest.raises(FileContractError, match="file.unsupported_format"):
        validate_upload_format(
            filename="a.pdf", media_type="application/pdf", content=b"not-pdf"
        )


def test_txt_docx_and_pdf_extract_text_and_tokens() -> None:
    """三种受支持格式都形成非空正文、版本和确定 Token 计数。"""

    txt = extract_file_text("你好 MetaHub".encode(), "txt")
    docx = extract_file_text(_docx_with_text(), "docx")
    pdf = extract_file_text(_pdf_with_text("Hello PDF"), "pdf")
    assert txt.text == "你好 MetaHub"
    assert txt.estimated_tokens == estimate_file_tokens(txt.text)
    assert "合同正文" in docx.text and "表格字段" in docx.text
    assert "Hello PDF" in pdf.text
    assert {txt.extractor_version, docx.extractor_version, pdf.extractor_version} == {
        "txt-v1",
        "docx-v1",
        "pdf-v1",
    }


def test_invalid_txt_empty_pdf_and_dangerous_docx_fail_closed() -> None:
    """无效编码、无文本 PDF 和 DOCX 嵌入对象不会被误标为成功。"""

    with pytest.raises(FileContractError, match="file.structure_invalid"):
        extract_file_text(b"\xff\xfe", "txt")
    with pytest.raises(FileContractError, match="file.no_extractable_text"):
        extract_file_text(_pdf_with_text("   "), "pdf")

    source = _docx_with_text()
    inbound = zipfile.ZipFile(io.BytesIO(source))
    target = io.BytesIO()
    with zipfile.ZipFile(target, "w") as outbound:
        for entry in inbound.infolist():
            outbound.writestr(entry, inbound.read(entry.filename))
        outbound.writestr("word/embeddings/object1.bin", b"opaque")
    inbound.close()
    with pytest.raises(FileContractError, match="file.structure_invalid"):
        extract_file_text(target.getvalue(), "docx")


def test_file_token_estimator_is_deterministic_for_cjk_other_and_whitespace() -> None:
    """Token v1 对 CJK、其他字符和空白使用冻结整数规则。"""

    assert estimate_file_tokens("中文") == 2
    assert estimate_file_tokens("abcd") == 1
    assert estimate_file_tokens("abcde") == 2
    assert estimate_file_tokens("中 abcd\n文") == 3
    assert estimate_file_tokens(" \r\n\t") == 0


def test_file_dlp_scans_full_text_and_counts_distinct_rules() -> None:
    """窗口后半段、跨窗 regex 和重复文本均只按不同规则数计数。"""

    prefix = "a" * (FILE_DLP_WINDOW_CODE_POINTS - 2)
    text = f"{prefix}SE" + "CRET and SECRET again"
    rules = (
        _rule("dlp_log", pattern="SECRET", action="log_only", order=0),
        _rule(
            "dlp_block",
            pattern="SE\\s*CRET",
            action="block",
            order=1,
            rule_type="regex",
        ),
    )
    decision = scan_file_text(text, rules)
    assert decision.status == "blocked"
    assert decision.action == "block"
    assert decision.effective_dlp_rule_id == "dlp_block"
    assert decision.match_count == 2
    assert len(decision.matches) == 2
    assert decision.ruleset_fingerprint.startswith("sha256:")


def test_file_cipher_is_random_owner_bound_and_repr_safe() -> None:
    """密文随机且跨 tenant/blob owner 交换认证失败，DTO repr 不含明文。"""

    secret = DEFAULT_FILE_ENCRYPTION_SECRET
    file_id = "file_" + "a" * 32
    content = b"confidential body"
    fingerprint = file_content_fingerprint(
        content,
        secret=secret,
        tenant_id="tenant_a",
        file_id=file_id,
        blob_kind="original",
    )
    owner = FileBlobOwner(
        tenant_id="tenant_a",
        file_id=file_id,
        blob_kind="original",
        content_fingerprint=fingerprint,
    )
    first = encrypt_file_blob(content, secret=secret, owner=owner)
    second = encrypt_file_blob(content, secret=secret, owner=owner)
    assert first != second
    assert content not in first
    assert decrypt_file_blob(first, secret=secret, owner=owner) == content
    assert "confidential" not in repr(owner)
    wrong_owner = FileBlobOwner(
        tenant_id="tenant_b",
        file_id=file_id,
        blob_kind="original",
        content_fingerprint=fingerprint,
    )
    with pytest.raises(FileContractError, match="file.integrity_failed"):
        decrypt_file_blob(first, secret=secret, owner=wrong_owner)


def test_file_secret_rejects_prod_default_and_cross_domain_reuse() -> None:
    """生产默认值和与其他密钥域复用都使文件能力失败关闭。"""

    assert validate_file_secret(
        DEFAULT_FILE_ENCRYPTION_SECRET,
        production=False,
    )
    with pytest.raises(FileContractError, match="file.storage_unavailable"):
        validate_file_secret(DEFAULT_FILE_ENCRYPTION_SECRET, production=True)
    with pytest.raises(FileContractError, match="file.storage_unavailable"):
        validate_file_secret(
            "x" * 40,
            production=False,
            forbidden_secrets=("x" * 40,),
        )


def test_encrypted_local_blob_store_publishes_reads_and_rejects_paths(
    tmp_path: Path,
) -> None:
    """本地 store 只接受固定对象键并以密文 staging 原子发布。"""

    store = EncryptedLocalBlobStore(tmp_path / "blob-root")
    file_id = "file_" + "b" * 32
    key = file_object_key(file_id, "original")
    staged = store.stage(b"ciphertext-only", object_key=key)
    assert "blob-root" not in repr(staged)
    assert store.publish(staged) == key
    assert store.read(key) == b"ciphertext-only"
    with pytest.raises(FileContractError, match="file.integrity_failed"):
        store.read("../outside")
    store.delete(key)
    store.delete(key)


def test_file_runtime_config_is_bounded_private_and_prod_optional(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """文件配置在 dev 可运行、prod 缺失只关闭能力，secret/root 不进入 dump/repr。"""

    monkeypatch.delenv("METAHUB_CONFIG_FILE", raising=False)
    monkeypatch.setenv("METAHUB_ENV", "dev")
    monkeypatch.delenv("METAHUB_FILE_ENCRYPTION_SECRET", raising=False)
    config = load_app_config("control-plane-api")
    assert config.file_encryption_secret == DEFAULT_FILE_ENCRYPTION_SECRET
    assert config.file_max_upload_bytes == 20 * 1024 * 1024
    assert config.file_parse_concurrency == 2
    assert "file_encryption_secret" not in config.model_dump()
    assert config.file_storage_root not in repr(config)

    monkeypatch.setenv("METAHUB_ENV", "prod")
    prod = load_app_config("control-plane-api")
    assert prod.file_encryption_secret is None
    assert prod.features["file_processing_enabled"] is False


def test_file_runtime_config_rejects_prod_default_reuse_and_bad_bounds(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """生产开发默认值、跨域复用和超界数值都在配置加载阶段拒绝。"""

    monkeypatch.delenv("METAHUB_CONFIG_FILE", raising=False)
    monkeypatch.setenv("METAHUB_ENV", "prod")
    monkeypatch.setenv("METAHUB_FILE_ENCRYPTION_SECRET", DEFAULT_FILE_ENCRYPTION_SECRET)
    with pytest.raises(ConfigError):
        load_app_config("control-plane-api")

    monkeypatch.setenv("METAHUB_ENV", "dev")
    monkeypatch.setenv("METAHUB_AUTH_TOKEN_SECRET", "a" * 40)
    monkeypatch.setenv("METAHUB_FILE_ENCRYPTION_SECRET", "a" * 40)
    with pytest.raises(ConfigError):
        load_app_config("control-plane-api")

    monkeypatch.setenv("METAHUB_FILE_ENCRYPTION_SECRET", "b" * 40)
    monkeypatch.setenv("METAHUB_FILE_PARSE_CONCURRENCY", "9")
    with pytest.raises(ConfigError):
        load_app_config("control-plane-api")
