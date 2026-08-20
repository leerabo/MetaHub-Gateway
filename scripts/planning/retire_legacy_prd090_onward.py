"""Retire legacy PRD-090..108 status pointers without rewriting audit rows."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
PLAN = ROOT / "企业级AI_Gateway_PRD开发总清单_给CodexClaude.xlsx"
SHEET = "12_开发状态"

OLD_A2 = (
    "Updated 2026-08-18: PRD-089 completed; PRD-090 is the next unfinished item; "
    "all_108_prds_complete=false."
)
OLD_B26 = (
    "Cross-checked next unfinished item: PRD-090. It is not coded, tested, migrated, "
    "or marked complete."
)
NEW_A2 = (
    "Updated 2026-08-20: PRD-001..089 completed; legacy PRD-090..108 retired and "
    "forbidden; all_108_prds_complete=false."
)
NEW_B26 = (
    "Legacy PRD-090..108 are retained only for audit and are not active work. The next "
    "item is defined only by 企业级AI_Gateway_SaaS增量开发计划_PRD090起.xlsx; "
    "all_108_prds_complete=false."
)


def main() -> None:
    workbook = load_workbook(PLAN)
    worksheet = workbook[SHEET]
    current = (worksheet["A2"].value, worksheet["B26"].value)
    allowed = {(OLD_A2, OLD_B26), (NEW_A2, NEW_B26)}
    if current not in allowed:
        raise RuntimeError("legacy status cells changed outside the controlled retirement contract")
    worksheet["A2"] = NEW_A2
    worksheet["B26"] = NEW_B26
    workbook.save(PLAN)

    verified = load_workbook(PLAN, read_only=True, data_only=False)[SHEET]
    if (verified["A2"].value, verified["B26"].value) != (NEW_A2, NEW_B26):
        raise RuntimeError("legacy status retirement verification failed")
    print("legacy_status=retired")
    print("retired_range=PRD-090..108")


if __name__ == "__main__":
    main()
