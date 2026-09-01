"""Validate the PRD-090 SaaS architecture and roadmap baseline."""

from __future__ import annotations

import argparse
import ast
import hashlib
import json
import re
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OLD_PLAN = ROOT / "企业级AI_Gateway_PRD开发总清单_给CodexClaude.xlsx"
NEW_PLAN = ROOT / "企业级AI_Gateway_SaaS增量开发计划_PRD090起.xlsx"
CONSOLE_SERVICE = ROOT / "control-plane-api/src/metahub_control/console/service.py"
README = ROOT / "README.md"
ROADMAP_DOC = ROOT / "docs/planning/MetaHub-SaaS-PRD-090-onward.md"
TEST_RECORD = ROOT / "docs/testing/PRD-090-tests.md"
FIRST_PRD = 90
OLD_RETIRED_LAST = 108

REQUIRED_DOCS = {
    ROOT / "docs/prd/PRD-090-saas-architecture-migration-baseline.md": (
        "ADR-090-01",
        "现状证据",
        "迁移矩阵",
        "64项导航处置",
        "迁移波次",
        "故障、回滚与兼容",
        "两轮编码前自审",
    ),
    ROOT / "docs/architecture/saas-target-architecture.md": (
        "部署形态",
        "目标拓扑",
        "身份与租户模型",
        "服务边界",
        "控制面到数据面合同",
        "统一可靠性目标",
    ),
    ROOT / "docs/security/saas-threat-model.md": (
        "STRIDE",
        "信任边界",
        "P0/P1 威胁与控制",
        "统一失败策略",
        "TM-01",
        "TM-14",
    ),
    ROOT / "docs/testing/PRD-090-tests.md": (
        "两轮自审",
        "空 SQLite 实际迁移",
        "十一类测试映射",
        "G0与完成状态",
        "PRD-091",
    ),
    ROOT / "docs/prd/PRD-092-self-service-registration-and-verification.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "METAHUB_IDENTITY_PII_ENCRYPTION_SECRET",
    ),
    ROOT / "docs/api/global-identity.md": (
        "POST | `/identity/registrations`",
        "METAHUB_IDENTITY_PII_ENCRYPTION_SECRET",
        "identity.unavailable",
    ),
    ROOT / "docs/testing/PRD-092-tests.md": (
        "状态：已完成",
        "敏感审计",
        "PRD-093",
    ),
    ROOT / "docs/prd/PRD-093-workspace-creation.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "workspace_user_projection_claims",
    ),
    ROOT / "docs/testing/PRD-093-tests.md": (
        "状态：已完成",
        "敏感审计",
        "PRD-094",
        "0085",
    ),
    ROOT / "docs/prd/PRD-094-membership-and-workspace-switching.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "METAHUB_WORKSPACE_SESSION_SECRET",
        "0086",
    ),
    ROOT / "docs/api/workspace-access.md": (
        "POST | `/identity/workspace-switches`",
        "GET | `/workspace/session`",
        "POST | `/membership-invitations`",
        "GET | `/console/navigation`",
        "METAHUB_WORKSPACE_SESSION_SECRET",
    ),
    ROOT / "docs/testing/PRD-094-tests.md": (
        "状态：已完成",
        "敏感审计",
        "PRD-095",
        "0086",
    ),
    ROOT / "docs/prd/PRD-095-platform-iam-controlled-support.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "0087",
    ),
    ROOT / "docs/api/platform-access.md": (
        "POST | `/platform/bootstrap`",
        "POST | `/platform/sessions`",
        "POST | `/platform/support-grants`",
        "next_cursor",
        "METAHUB_PLATFORM_SESSION_SECRET",
        "GET | `/platform/workspaces`",
        "POST | `/platform/workspaces/{workspace_id}/suspend`",
        "POST | `/platform/workspaces/{workspace_id}/owner-notifications`",
        "Idempotency-Key",
        "transition_pending",
    ),
    ROOT / "docs/testing/PRD-095-tests.md": (
        "状态：已完成",
        "敏感审计",
        "PRD-096",
        "0087",
    ),
    ROOT / "docs/prd/PRD-096-platform-workspace-directory-operations.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "0088",
    ),
    ROOT / "docs/testing/PRD-096-tests.md": (
        "状态：已完成",
        "敏感审计",
        "PRD-097",
        "0088",
    ),
    ROOT / "docs/prd/PRD-119-official-provider-catalog-regional-compliance.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "0111",
    ),
    ROOT / "docs/testing/PRD-119-tests.md": (
        "状态：已完成",
        "敏感",
        "PRD-120",
        "0111",
    ),
    ROOT / "docs/prd/PRD-120-edge-agent-device-identity-mtls.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "0112",
    ),
    ROOT / "docs/testing/PRD-120-tests.md": (
        "状态：已完成",
        "敏感",
        "PRD-121",
        "0112",
    ),
    ROOT / "docs/prd/PRD-121-edge-desired-state-signed-snapshot-offline-lkg.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "0113",
    ),
    ROOT / "docs/testing/PRD-121-tests.md": (
        "状态：已完成",
        "敏感",
        "PRD-122",
        "0113",
    ),
    ROOT / "docs/prd/PRD-122-private-gateway-telemetry-usage-failure-sync.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "0114",
    ),
    ROOT / "docs/testing/PRD-122-tests.md": (
        "状态：已完成",
        "敏感",
        "PRD-123",
        "0114",
    ),
    ROOT / "docs/prd/PRD-123-private-gateway-local-secret-custody-rotation.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "0115",
    ),
    ROOT / "docs/testing/PRD-123-tests.md": (
        "状态：已完成",
        "敏感",
        "PRD-124",
        "0115",
    ),
    ROOT / "docs/prd/PRD-124-private-gateway-installer-preflight-onboarding.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "0116",
    ),
    ROOT / "docs/testing/PRD-124-tests.md": (
        "状态：已完成",
        "敏感",
        "PRD-125",
        "0116",
    ),
    ROOT / "docs/prd/PRD-125-private-deployment-validation-redacted-support-bundle.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "0117",
    ),
    ROOT / "docs/testing/PRD-125-tests.md": (
        "状态：已完成",
        "敏感",
        "PRD-126",
        "0117",
    ),
    ROOT / "docs/prd/PRD-126-fleet-assets-health-controlled-diagnostics.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "0118",
    ),
    ROOT / "docs/testing/PRD-126-tests.md": (
        "状态：已完成",
        "敏感",
        "PRD-127",
        "0118",
    ),
    ROOT / "docs/prd/PRD-127-fleet-staged-upgrade-rollback-maintenance-window.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "0119",
    ),
    ROOT / "docs/testing/PRD-127-tests.md": (
        "状态：已完成",
        "敏感",
        "PRD-128",
        "0119",
    ),
    ROOT / "docs/prd/PRD-128-private-gateway-license-capacity-offline-grace.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "0120",
    ),
    ROOT / "docs/testing/PRD-128-tests.md": (
        "状态：已完成",
        "敏感",
        "PRD-129",
        "0120",
    ),
    ROOT / "docs/prd/PRD-129-private-gateway-backup-restore-dr-drill.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "0121",
    ),
    ROOT / "docs/testing/PRD-129-tests.md": (
        "状态：已完成",
        "敏感",
        "PRD-130",
        "0121",
    ),
    ROOT / "docs/prd/PRD-130-isolated-environment-signed-bundle-conflict-reconciliation.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "0122",
    ),
    ROOT / "docs/testing/PRD-130-tests.md": (
        "状态：已完成",
        "敏感",
        "PRD-131",
        "0122",
    ),
    ROOT / "docs/prd/PRD-131-legacy-tenant-direct-private-migration-cutover.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "0123",
    ),
    ROOT / "docs/testing/PRD-131-tests.md": (
        "状态：已完成",
        "敏感",
        "PRD-132",
        "0123",
    ),
    ROOT / "docs/prd/PRD-132-customer-ticket-sla-secure-diagnostics.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "0124",
    ),
    ROOT / "docs/testing/PRD-132-tests.md": (
        "状态：已完成",
        "敏感",
        "PRD-133",
        "0124",
    ),
    ROOT / "docs/prd/PRD-133-public-control-api-sdk-version-deprecation.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "0125",
    ),
    ROOT / "docs/testing/PRD-133-tests.md": (
        "状态：已完成",
        "敏感",
        "PRD-134",
        "0125",
    ),
    ROOT / "docs/prd/PRD-134-global-account-privacy-consent-data-subject-request.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "0126",
    ),
    ROOT / "docs/testing/PRD-134-tests.md": (
        "状态：已完成",
        "敏感",
        "PRD-135",
        "0126",
    ),
    ROOT / "docs/prd/PRD-135-compliance-evidence-center.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "0127",
    ),
    ROOT / "docs/testing/PRD-135-tests.md": (
        "状态：已完成",
        "敏感",
        "PRD-136",
        "0127",
    ),
    ROOT / "docs/prd/PRD-136-sbom-artifact-signing-vulnerability-provenance.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "0128",
    ),
    ROOT / "docs/testing/PRD-136-tests.md": (
        "状态：已完成",
        "敏感",
        "PRD-137",
        "0128",
    ),
    ROOT / "docs/prd/PRD-137-tenant-isolation-continuous-security-validation.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "0129",
    ),
    ROOT / "docs/testing/PRD-137-tests.md": (
        "状态：已完成",
        "敏感",
        "PRD-138",
        "0129",
    ),
    ROOT / "docs/prd/PRD-138-saas-release-candidate-launch-gates.md": (
        "状态：已完成",
        "第一轮自审",
        "第二轮自审",
        "0130",
        "release_candidate",
    ),
    ROOT / "docs/testing/PRD-138-tests.md": (
        "状态：已完成",
        "敏感",
        "PRD-139",
        "0130",
        "release_candidate",
    ),
    ROOT / "docs/prd/PRD-139-channel-partner-attribution-management.md": (
        "状态：已完成",
        "第一轮：",
        "第二轮：",
        "0131",
    ),
    ROOT / "docs/testing/PRD-139-tests.md": (
        "状态：已完成",
        "敏感",
        "PRD-140",
        "0131",
    ),
    ROOT / "docs/prd/PRD-140-saas-platform-customer-success-renewal.md": (
        "状态：已完成",
        "第一轮：",
        "第二轮：",
        "0132",
    ),
    ROOT / "docs/testing/PRD-140-tests.md": (
        "状态：已完成",
        "敏感",
        "PRD-141",
        "0132",
    ),
    ROOT / "docs/prd/PRD-141-saas-platform-growth-funnel-campaign-attribution.md": (
        "状态：已完成",
        "第一轮：",
        "第二轮：",
        "0133",
    ),
    ROOT / "docs/testing/PRD-141-tests.md": (
        "状态：已完成",
        "敏感",
        "PRD-142",
        "0133",
    ),
    ROOT / "docs/prd/PRD-142-platform-channel-deployment-branding-domain.md": (
        "状态：已完成",
        "第一轮：",
        "第二轮：",
        "0134",
    ),
    ROOT / "docs/testing/PRD-142-tests.md": (
        "状态：已完成",
        "敏感",
        "PRD-143",
        "0134",
    ),
}

REQUIRED_MASTER_FIELDS = (
    "PRD_ID",
    "计划状态",
    "阶段",
    "优先级",
    "能力层",
    "模块",
    "PRD名称",
    "主要用户/角色",
    "用户故事",
    "需求范围/功能说明",
    "明确不做/边界",
    "页面/入口",
    "API/动作",
    "核心数据对象",
    "关键权限",
    "验收标准",
    "依赖关系",
    "数据迁移/回滚",
    "安全与合规",
    "测试重点",
    "性能/SLO",
    "预计工作量",
    "建议团队",
    "交付证据",
)

SECRET_PATTERNS = {
    "aws_access_key": re.compile(r"AKIA[0-9A-Z]{16}"),
    "provider_key": re.compile(r"\bsk-[A-Za-z0-9_-]{20,}\b"),
    "bearer_token": re.compile(r"Bearer\s+[A-Za-z0-9._-]{24,}", re.IGNORECASE),
    "private_key": re.compile(r"-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----"),
}

POST_NAVIGATION_BASELINE_IDS = frozenset({"migration-operations", "support-tickets"})
RETIRED_SOURCE_NAVIGATION_IDS = frozenset({"branding", "growth-analytics"})

ROLE_ALIAS_PATTERN = re.compile(
    r"(?<![A-Za-z0-9_])(?:ops|security|support|billing|finance|verification_reviewer)(?![A-Za-z0-9_])"
)
ROLE_ID_PATTERN = re.compile(
    r"\b(?:platform_[a-z_]+|[a-z]+_admin|release_manager|owner|admin|developer|auditor|viewer|service_account)\b"
)


@dataclass
class ValidationReport:
    checks: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    next_prd: str | None = None

    def check(self, condition: bool, message: str) -> None:
        if condition:
            self.checks.append(message)
        else:
            self.errors.append(message)

    @property
    def ok(self) -> bool:
        return not self.errors


def _rows_by_header(
    path: Path,
    sheet_name: str,
    *,
    header_key: str | None = None,
) -> tuple[list[str], list[dict[str, object]]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    worksheet = workbook[sheet_name]
    raw_rows = list(worksheet.iter_rows(values_only=True))
    header_index = 0
    if header_key is not None:
        header_index = next(
            index for index, row in enumerate(raw_rows[:5]) if header_key in row
        )
    headers = [str(value) for value in raw_rows[header_index]]
    return headers, [
        dict(zip(headers, row, strict=True))
        for row in raw_rows[header_index + 1 :]
        if any(value is not None for value in row)
    ]


def _explicit_dependencies(value: object) -> set[int]:
    return {int(match) for match in re.findall(r"(?<!\d)(\d{3})(?!\d)", str(value or ""))}


def _continuous_prd_ids(values: list[str], *, first: int = FIRST_PRD) -> bool:
    numbers: list[int] = []
    for value in values:
        match = re.fullmatch(r"PRD-(\d{3,})", value)
        if match is None:
            return False
        numbers.append(int(match.group(1)))
    return bool(numbers) and numbers == list(range(first, first + len(numbers)))


def _workbook_text(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=False)
    return "\n".join(
        str(value)
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows(values_only=True)
        for value in row
        if value is not None
    )


def _document_texts(paths: Iterable[Path]) -> str:
    return "\n".join(path.read_text(encoding="utf-8") for path in paths if path.exists())


def _xlsx_package_is_safe(path: Path) -> bool:
    with zipfile.ZipFile(path) as archive:
        names = {name.lower() for name in archive.namelist()}
    return not any(
        name.startswith("xl/externallinks/") or name.endswith("vbaproject.bin")
        for name in names
    )


def _console_navigation_ids() -> list[str]:
    tree = ast.parse(CONSOLE_SERVICE.read_text(encoding="utf-8"))
    assignment = next(
        node
        for node in tree.body
        if isinstance(node, ast.AnnAssign)
        and isinstance(node.target, ast.Name)
        and node.target.id == "NAVIGATION_DEFINITIONS"
    )
    assert isinstance(assignment.value, ast.Tuple)
    return [str(ast.literal_eval(element.args[0])) for element in assignment.value.elts]


def _source_navigation_ids() -> list[str]:
    current = [
        item
        for item in _console_navigation_ids()
        if item not in POST_NAVIGATION_BASELINE_IDS
    ]
    return [*current, *sorted(RETIRED_SOURCE_NAVIGATION_IDS - set(current))]


def validate(*, expected_next: str | None = None) -> ValidationReport:
    report = ValidationReport()
    report.check(OLD_PLAN.is_file(), "old roadmap retained for audit")
    report.check(NEW_PLAN.is_file(), "authoritative SaaS roadmap exists")
    if not OLD_PLAN.is_file() or not NEW_PLAN.is_file():
        return report

    headers, rows = _rows_by_header(NEW_PLAN, "01_PRD总表")
    report.check(all(field in headers for field in REQUIRED_MASTER_FIELDS), "master sheet has required fields")
    actual_ids = [str(row["PRD_ID"]) for row in rows]
    report.check(
        _continuous_prd_ids(actual_ids),
        "active PRD identifiers start at PRD-090 and remain continuous without a fixed upper limit",
    )
    current_range = f"{actual_ids[0]} 至 {actual_ids[-1]}" if actual_ids else ""
    readme_text = README.read_text(encoding="utf-8") if README.is_file() else ""
    roadmap_text = ROADMAP_DOC.read_text(encoding="utf-8") if ROADMAP_DOC.is_file() else ""
    test_record_text = TEST_RECORD.read_text(encoding="utf-8") if TEST_RECORD.is_file() else ""
    report.check(
        bool(current_range) and current_range in readme_text and current_range in roadmap_text,
        "README and roadmap document match the workbook's current open-ended PRD range",
    )
    report.check(
        f"当前已细化{len(actual_ids)}条" in test_record_text,
        "PRD-090 test record matches the workbook's current planned item count",
    )
    report.check(
        "不限定需求总数或最终编号" in readme_text
        and "不限定需求总数或最终编号" in roadmap_text,
        "README and roadmap document do not declare a fixed PRD ceiling",
    )
    titles = [str(row["PRD名称"]) for row in rows]
    report.check(len(titles) == len(set(titles)), "roadmap titles are unique")

    workbook_structure = load_workbook(NEW_PLAN, read_only=True, data_only=False)
    report.check(all(sheet.sheet_state == "visible" for sheet in workbook_structure.worksheets), "roadmap has no hidden sheets")
    report.check(
        all(cell.data_type != "f" for sheet in workbook_structure.worksheets for row in sheet.iter_rows() for cell in row),
        "roadmap has no formulas",
    )
    report.check(_xlsx_package_is_safe(NEW_PLAN), "roadmap has no macros or external links")

    missing = [
        f"{row['PRD_ID']}:{field}"
        for row in rows
        for field in REQUIRED_MASTER_FIELDS
        if row.get(field) is None or not str(row[field]).strip()
    ]
    report.check(not missing, "all PRDs contain scope, boundary, permission, acceptance, migration, security, tests and SLO")
    if missing:
        report.errors.append("missing fields: " + ", ".join(missing[:20]))

    rows_by_id = {str(row["PRD_ID"]): row for row in rows}
    prd091 = rows_by_id.get("PRD-091", {})
    report.check(
        "WorkspaceMember" not in str(prd091.get("需求范围/功能说明") or "")
        and "platform_" not in str(prd091.get("关键权限") or "")
        and "/identity/credentials/password/rotate" in str(prd091.get("API/动作") or "")
        and "identity.credential_rotated" in str(prd091.get("事件/Webhook") or ""),
        "PRD-091 remains global-account-only and matches its delivered credential contract",
    )
    prd092 = rows_by_id.get("PRD-092", {})
    prd092_apis = str(prd092.get("API/动作") or "")
    prd092_security = str(prd092.get("安全与合规") or "")
    report.check(
        all(
            endpoint in prd092_apis
            for endpoint in (
                "/identity/registrations",
                "/identity/verifications",
                "/identity/verifications/resend",
                "/identity/password-resets",
                "/identity/password-resets/confirm",
            )
        )
        and "AES-256-GCM" in prd092_security
        and "Challenge根密钥" in prd092_security
        and "PII根密钥" in prd092_security,
        "PRD-092 matches its delivered self-service and independent secret-domain contracts",
    )
    prd093 = rows_by_id.get("PRD-093", {})
    prd093_apis = str(prd093.get("API/动作") or "")
    prd093_migration = str(prd093.get("数据迁移/回滚") or "")
    prd093_security = str(prd093.get("安全与合规") or "")
    report.check(
        all(
            endpoint in prd093_apis
            for endpoint in (
                "POST /workspaces",
                "GET /workspaces",
                "GET /workspaces/{workspace_id}",
                "PATCH /workspaces/{workspace_id}",
            )
        )
        and "0085" in prd093_migration
        and "WorkspaceTenantBinding" in prd093_migration
        and "GlobalSession" in prd093_security
        and "Idempotency-Key" in prd093_security,
        "PRD-093 matches its delivered Workspace, migration and request-integrity contracts",
    )
    prd094 = rows_by_id.get("PRD-094", {})
    prd094_apis = str(prd094.get("API/动作") or "")
    prd094_objects = str(prd094.get("核心数据对象") or "")
    prd094_migration = str(prd094.get("数据迁移/回滚") or "")
    prd094_security = str(prd094.get("安全与合规") or "")
    prd094_boundary = str(prd094.get("明确不做/边界") or "")
    report.check(
        all(
            endpoint in prd094_apis
            for endpoint in (
                "GET /identity/workspaces",
                "POST /identity/workspace-switches",
                "GET /workspace/session",
                "GET /memberships",
                "PATCH /memberships/{membership_id}",
                "POST /membership-invitations",
                "POST /workspace-owner-transfers",
                "GET /console/navigation",
            )
        )
        and all(
            object_name in prd094_objects
            for object_name in (
                "WorkspaceMembership",
                "WorkspaceSession",
                "MembershipInvitation",
                "OwnerTransfer",
                "MembershipEvent",
                "MembershipEventOutbox",
            )
        )
        and "0086" in prd094_migration
        and "pristine" in prd094_migration
        and "METAHUB_WORKSPACE_SESSION_SECRET" in prd094_security
        and "每请求" in prd094_security
        and all(
            subject in prd094_boundary
            for subject in (
                "GlobalSession",
                "Legacy Bearer",
                "Platform",
                "Device",
                "service_account",
            )
        ),
        "PRD-094 matches its delivered Workspace access, migration and subject-isolation contracts",
    )
    prd095 = rows_by_id.get("PRD-095", {})
    prd095_apis = str(prd095.get("API/动作") or "")
    prd095_objects = str(prd095.get("核心数据对象") or "")
    prd095_migration = str(prd095.get("数据迁移/回滚") or "")
    prd095_security = str(prd095.get("安全与合规") or "")
    prd095_tests = str(prd095.get("测试重点") or "")
    report.check(
        all(
            endpoint in prd095_apis
            for endpoint in (
                "POST /identity/mfa/totp/enrollments",
                "POST /platform/bootstrap",
                "POST /platform/sessions",
                "POST /platform/support-grants",
                "POST /workspace/support-grants/{grant_id}/approve",
                "GET /support/workspace/diagnostics",
                "POST /platform/break-glass-sessions",
                "GET /platform/audit-logs",
            )
        )
        and all(
            object_name in prd095_objects
            for object_name in (
                "AccountMfaFactor",
                "PlatformSession",
                "SupportGrantApproval",
                "SupportSession",
                "BreakGlassReview",
                "PlatformAuditLog",
            )
        )
        and "0087" in prd095_migration
        and "pristine" in prd095_migration
        and "Global/Workspace/Platform/Support" in prd095_security
        and "100并发审批/Session" in prd095_tests,
        "PRD-095 matches its delivered Platform IAM, support and subject-isolation contracts",
    )
    prd096 = rows_by_id.get("PRD-096", {})
    prd096_apis = str(prd096.get("API/动作") or "")
    prd096_objects = str(prd096.get("核心数据对象") or "")
    prd096_migration = str(prd096.get("数据迁移/回滚") or "")
    prd096_security = str(prd096.get("安全与合规") or "")
    prd096_tests = str(prd096.get("测试重点") or "")
    report.check(
        all(
            endpoint in prd096_apis
            for endpoint in (
                "GET /platform/workspaces",
                "GET /platform/workspaces/{workspace_id}/health",
                "POST /platform/workspaces/{workspace_id}/suspend",
                "POST/DELETE /platform/workspaces/{workspace_id}/risk-tags",
                "POST /platform/workspaces/{workspace_id}/owner-notifications",
                "GET /platform/workspaces/{workspace_id}/events",
            )
        )
        and all(
            object_name in prd096_objects
            for object_name in (
                "WorkspaceOperationalState",
                "WorkspaceRiskTag",
                "WorkspaceHealthProjection",
                "PlatformActionIntent",
                "WorkspaceOwnerNotificationOutbox",
                "WorkspaceOperationalEvent",
            )
        )
        and "0088" in prd096_migration
        and "pristine" in prd096_migration
        and "HMAC游标" in prd096_security
        and "五角色字段裁剪" in prd096_tests,
        "PRD-096 matches its delivered directory, lifecycle and privacy contracts",
    )
    prd106 = rows_by_id.get("PRD-106", {})
    report.check(
        "租户管理员" not in str(prd106.get("主要用户/角色") or "")
        and "不共用平台工作台、API或授权" in str(prd106.get("需求范围/功能说明") or "")
        and "/platform/jobs/" in str(prd106.get("API/动作") or ""),
        "platform scheduler and enterprise batch jobs have separate workbenches and APIs",
    )
    prd142 = rows_by_id.get("PRD-142", {})
    report.check(
        "独立导航、DTO、写权限和owner" in str(prd142.get("需求范围/功能说明") or ""),
        "platform branding and private Deployment appearance are explicitly separated",
    )
    prd143 = rows_by_id.get("PRD-143", {})
    report.check(
        "Workspace" not in str(prd143.get("主要用户/角色") or "")
        and "企业开发者只在Provider详情" in str(prd143.get("需求范围/功能说明") or ""),
        "Provider engineering is platform-only while enterprise BYOK gets a separate projection",
    )
    prd146_scope = str(rows_by_id.get("PRD-146", {}).get("需求范围/功能说明") or "")
    report.check(
        all(label in prd146_scope for label in ("BYOK", "企业内部财务", "文件安全", "异步批处理", "MCP")),
        "enterprise Gateway scope explicitly retains AI governance, files, batch and MCP",
    )

    seen_unstarted = False
    status_order_ok = True
    for row in rows:
        status = str(row["计划状态"])
        if status == "未启动":
            seen_unstarted = True
            report.next_prd = report.next_prd or str(row["PRD_ID"])
        elif status == "已完成" and not seen_unstarted:
            continue
        else:
            status_order_ok = False
    report.check(status_order_ok, "completed PRDs form a contiguous prefix")
    if expected_next is not None:
        report.check(report.next_prd == expected_next, f"next PRD is {expected_next}")

    dependency_ok = True
    graph: dict[int, set[int]] = {}
    for row in rows:
        current = int(str(row["PRD_ID"]).removeprefix("PRD-"))
        dependencies = _explicit_dependencies(row["依赖关系"])
        graph[current] = {dependency for dependency in dependencies if dependency >= FIRST_PRD}
        if any(dependency >= current for dependency in dependencies):
            dependency_ok = False
    report.check(dependency_ok, "explicit PRD dependencies only point backward")
    report.check(_is_acyclic(graph), "explicit PRD dependency graph is acyclic")

    old_headers, old_rows = _rows_by_header(OLD_PLAN, "01_PRD总表")
    report.check("PRD_ID" in old_headers and "PRD名称" in old_headers, "old roadmap has auditable PRD columns")
    old_retired = {
        str(row["PRD_ID"]): str(row["PRD名称"])
        for row in old_rows
        if f"PRD-{FIRST_PRD:03d}" <= str(row["PRD_ID"]) <= f"PRD-{OLD_RETIRED_LAST:03d}"
    }
    legacy_headers, legacy_rows = _rows_by_header(NEW_PLAN, "07_旧计划废弃映射")
    report.check("旧计划状态" in legacy_headers, "retired mapping has explicit status")
    mapped = {str(row["PRD_ID"]): str(row["旧需求名称"]) for row in legacy_rows}
    report.check(mapped == old_retired, "old PRD-090..108 titles match retired mapping")
    report.check(
        len(legacy_rows) == 19 and all(row["旧计划状态"] == "已废弃-禁止开发" for row in legacy_rows),
        "all 19 old PRDs are retired and forbidden",
    )
    report.check(not (set(old_retired.values()) & set(titles)), "no retired title is active in the new master")

    workbook_text = _workbook_text(NEW_PLAN)
    _, permission_rows = _rows_by_header(NEW_PLAN, "05_角色权限矩阵", header_key="角色")
    platform_roles = {str(row["角色"]) for row in permission_rows if row.get("角色域") == "平台"}
    tenant_roles = {str(row["角色"]) for row in permission_rows if row.get("角色域") == "租户"}
    report.check(not (platform_roles & tenant_roles), "platform and tenant role codes do not overlap")
    report.check(
        {"platform_owner", "platform_channel"} <= platform_roles and {"owner", "admin", "auditor"} <= tenant_roles,
        "platform, channel and enterprise functional role namespaces are explicit",
    )

    role_contract_text = "\n".join(
        str(row.get(field) or "")
        for row in rows
        for field in ("主要用户/角色", "需求范围/功能说明", "关键权限")
    )
    report.check(ROLE_ALIAS_PATTERN.search(role_contract_text) is None, "role contracts use exact role ids, not aliases")
    referenced_roles = set(ROLE_ID_PATTERN.findall(role_contract_text))
    report.check(
        referenced_roles <= platform_roles | tenant_roles,
        "all explicit role ids are declared in the role matrix",
    )

    page_headers, page_rows = _rows_by_header(NEW_PLAN, "10_角色页面矩阵")
    required_page_fields = {"角色", "角色域", "允许导航/页面", "关键动作", "明确禁止", "分配主体"}
    report.check(required_page_fields <= set(page_headers), "role page matrix has required fields")
    page_roles = {str(row["角色"]) for row in page_rows}
    report.check(page_roles == platform_roles | tenant_roles, "every declared role has one page-access contract")
    role_domains = {str(row["角色"]): str(row["角色域"]) for row in permission_rows}
    report.check(
        len(page_rows) == len(page_roles)
        and all(str(row.get("角色域")) == role_domains.get(str(row.get("角色"))) for row in page_rows),
        "role page contracts are unique and remain in their IAM domain",
    )
    report.check(
        all(all(str(row.get(field) or "").strip() for field in required_page_fields) for row in page_rows),
        "role page contracts define pages, actions, prohibitions and assignment authority",
    )
    page_text = "\n".join(str(value) for row in page_rows for value in row.values())
    report.check(
        "企业owner/admin添加职能用户" in workbook_text or "owner/admin在自身委派上限内" in page_text,
        "enterprise administrators assign functional users within a delegation ceiling",
    )
    report.check(
        "直达URL" in workbook_text and "逐请求鉴权" in workbook_text,
        "server authorization remains authoritative beyond navigation visibility",
    )
    tenant_allowed_pages = "\n".join(
        str(row.get("允许导航/页面") or "") for row in page_rows if row.get("角色域") == "租户"
    )
    report.check(
        not any(label in tenant_allowed_pages for label in ("渠道", "客户管理", "增长分析")),
        "channel, customer-management and growth pages remain platform-only",
    )
    report.check(
        not any(
            label in tenant_allowed_pages
            for label in ("平台品牌", "Provider Adapter", "合同CI", "套餐目录", "争议裁决")
        ),
        "platform branding, Provider engineering and commercial decision pages remain platform-only",
    )
    report.check(
        all(label in tenant_allowed_pages for label in ("文件", "批处理", "MCP")),
        "enterprise role pages retain files, asynchronous batch and MCP capabilities",
    )
    page_by_role = {str(row.get("角色")): row for row in page_rows}
    platform_ops_pages = str(page_by_role.get("platform_ops", {}).get("允许导航/页面") or "")
    platform_ops_actions = str(page_by_role.get("platform_ops", {}).get("关键动作") or "")
    provider_engineer_pages = str(
        page_by_role.get("platform_provider_engineer", {}).get("允许导航/页面") or ""
    )
    report.check(
        "Provider工程" not in platform_ops_pages
        and "合同测试" not in platform_ops_actions
        and "合同CI" in provider_engineer_pages,
        "Provider engineering belongs to platform_provider_engineer, not platform_ops",
    )

    navigation_headers, navigation_rows = _rows_by_header(NEW_PLAN, "11_现有导航处置")
    required_navigation_fields = {
        "现有分组", "navigation_id", "现有页面", "处置", "目标产品域/入口", "企业可见角色",
        "平台可见角色", "现有API/数据/Scheduler证据", "承接PRD", "切换验收边界",
    }
    report.check(required_navigation_fields <= set(navigation_headers), "navigation disposition has required fields")
    console_navigation_ids = _console_navigation_ids()
    source_navigation_ids = _source_navigation_ids()
    disposition_ids = [str(row.get("navigation_id")) for row in navigation_rows]
    report.check(
        POST_NAVIGATION_BASELINE_IDS <= set(console_navigation_ids)
        and not (POST_NAVIGATION_BASELINE_IDS & set(source_navigation_ids))
        and RETIRED_SOURCE_NAVIGATION_IDS.isdisjoint(console_navigation_ids)
        and RETIRED_SOURCE_NAVIGATION_IDS <= set(source_navigation_ids),
        "post-baseline PRD-131 navigation remains explicit without rewriting the 64-item inventory",
    )
    report.check(
        len(source_navigation_ids) == 64
        and len(disposition_ids) == 64
        and len(set(disposition_ids)) == 64
        and set(disposition_ids) == set(source_navigation_ids),
        "all 64 source navigation items have exactly one disposition",
    )
    report.check(
        all(all(str(row.get(field) or "").strip() for field in required_navigation_fields) for row in navigation_rows),
        "every navigation disposition defines roles, implementation evidence, target PRD and cutover boundary",
    )
    disposition_by_id = {str(row.get("navigation_id")): row for row in navigation_rows}
    saas_only_ids = {
        "growth-analytics",
        "customer-health",
        "partner-portal",
        "adapter-specs",
        "adapter-contract-runs",
    }
    split_domain_ids = {
        "overview",
        "admin-assistant",
        "approvals",
        "branding",
        "provider-health",
        "error-maps",
        "model-prices",
        "plans",
        "billing-disputes",
        "policy-templates",
        "system-status",
        "config-summary",
    }
    global_shell_ids = {"language-settings"}
    retired_product_ui_ids = {"competitor-comparison"}
    enterprise_only_ids = set(source_navigation_ids) - (
        saas_only_ids | split_domain_ids | global_shell_ids | retired_product_ui_ids
    )
    report.check(
        len(enterprise_only_ids) == 45
        and len(split_domain_ids) == 12
        and len(saas_only_ids) == 5
        and len(global_shell_ids) == 1
        and len(retired_product_ui_ids) == 1,
        "navigation ownership closes as 45 enterprise, 12 split, 5 SaaS, 1 global and 1 retired",
    )
    report.check(
        all(
            str(disposition_by_id[item].get("企业可见角色")) == "无"
            and str(disposition_by_id[item].get("平台可见角色")) != "无"
            and str(disposition_by_id[item].get("处置")) == "迁出企业"
            for item in saas_only_ids
        ),
        "all five SaaS-only pages leave the enterprise navigation and retain platform owners",
    )
    report.check(
        all(
            str(disposition_by_id[item].get("企业可见角色")) != "无"
            and str(disposition_by_id[item].get("平台可见角色")) != "无"
            and "拆分" in str(disposition_by_id[item].get("处置"))
            for item in split_domain_ids
        ),
        "all twelve shared capabilities require separate platform and enterprise projections",
    )
    report.check(
        all(
            str(disposition_by_id[item].get("处置")).startswith("企业保留")
            and str(disposition_by_id[item].get("企业可见角色")) != "无"
            for item in enterprise_only_ids
        ),
        "all forty-five enterprise Gateway pages remain enterprise-owned",
    )
    report.check(
        disposition_by_id["language-settings"].get("处置") == "迁到全局账号壳层"
        and disposition_by_id["competitor-comparison"].get("处置") == "迁出产品控制台",
        "global preference and retired comparison content are not misclassified as SaaS operations",
    )
    report.check(
        all(
            str(disposition_by_id[item].get("承接PRD") or "").startswith("PRD-")
            for item in disposition_ids
        ),
        "every frozen navigation item is assigned to an explicit migration PRD",
    )

    object_headers, object_rows = _rows_by_header(NEW_PLAN, "03_核心数据对象")
    report.check("数据对象" in object_headers and "归属域" in object_headers, "object ownership sheet is structured")
    global_account_rows = [row for row in object_rows if row.get("数据对象") == "GlobalAccount"]
    report.check(
        bool(global_account_rows) and all(row.get("归属域") == "平台全局" for row in global_account_rows),
        "GlobalAccount is consistently classified as platform-global",
    )
    abuse_counter_rows = [row for row in object_rows if row.get("数据对象") == "AbuseCounter"]
    report.check(
        bool(abuse_counter_rows)
        and all(row.get("归属域") == "平台全局/Identity abuse domain" for row in abuse_counter_rows)
        and all("不带tenant_id" in str(row.get("租户字段规则") or "") for row in abuse_counter_rows),
        "AbuseCounter is owned by the tenant-free Global Identity abuse domain",
    )

    report.check("saas_incremental_roadmap_complete=false" in workbook_text, "new roadmap completion flag is false")
    report.check("all_108_prds_complete=false" in workbook_text, "legacy compatibility completion flag is false")
    report.check("all_108_prds_complete=true" not in workbook_text, "legacy true flag is absent")
    report.check("智谱" in workbook_text and "DeepSeek" in workbook_text, "default official provider boundary is present")
    report.check("OpenRouter" in workbook_text and "竞品" in workbook_text, "competitor boundary is present")
    report.check(
        "个人SaaS、企业SaaS、云管私有Gateway" in workbook_text
        and "完全离线仅支持签名包异步管理" in workbook_text,
        "three standard deployment forms and offline asynchronous capability are distinct",
    )
    report.check(
        "不设需求总数或最终编号上限" in workbook_text,
        "roadmap has no fixed PRD count or final identifier",
    )

    old_workbook = load_workbook(OLD_PLAN, read_only=True, data_only=False)
    old_status = old_workbook["12_开发状态"]
    report.check(
        "legacy PRD-090..108 retired and forbidden" in str(old_status["A2"].value)
        and "retained only for audit" in str(old_status["B26"].value),
        "legacy active status pointers are retired",
    )

    for path, markers in REQUIRED_DOCS.items():
        report.check(path.is_file(), f"required document exists: {path.relative_to(ROOT)}")
        if path.is_file():
            content = path.read_text(encoding="utf-8")
            report.check(all(marker in content for marker in markers), f"required sections exist: {path.name}")

    prd_text = (ROOT / "docs/prd/PRD-090-saas-architecture-migration-baseline.md").read_text(encoding="utf-8")
    report.check("待独立审查结论合入" not in prd_text, "both independent self-reviews are closed")
    report.check(prd_text.count("P0/P1已关闭") == 2, "both self-reviews explicitly close P0/P1 findings")

    architecture_text = (ROOT / "docs/architecture/saas-target-architecture.md").read_text(encoding="utf-8")
    threat_text = (ROOT / "docs/security/saas-threat-model.md").read_text(encoding="utf-8")
    security_contracts = {
        "five disjoint subjects": "五类issuer、audience和key domain不得复用",
        "immutable workspace binding": "WorkspaceTenantBinding(workspace_id UNIQUE, tenant_id UNIQUE",
        "edge CA and CSR proof": "独立Edge CA签发短期证书",
        "canonical signed envelope": "完整RFC 8785 envelope",
        "persistent generation watermark": "高水位跨重启/备份恢复持久且不得回拨",
        "offline revocation window": "trust-manifest TTL约束",
        "secret AAD owner binding": "execution_plane, workspace_id, deployment_id-or-cloud, secret_id, version, purpose",
        "data residency matrix": "数据类别与位置矩阵",
        "support dual-domain approval": "至少一位为该Workspace客户授权者、一位为平台授权者",
        "mixed credential rejection": "Bearer与Cookie",
    }
    for name, marker in security_contracts.items():
        report.check(marker in architecture_text, f"security contract exists: {name}")
    report.check("Owner与驻留验证矩阵" in threat_text, "threat model carries owner and residency controls")

    migration_versions = sorted(
        int(match.group(1))
        for path in (ROOT / "infra/migrations").glob("*.sql")
        if ".down." not in path.name and (match := re.match(r"(\d{4})_", path.name))
    )
    report.check(
        bool(migration_versions)
        and migration_versions[-1] == 134
        and (ROOT / "infra/migrations/0083_global_account_identity.sql").is_file()
        and (ROOT / "infra/migrations/0083_global_account_identity.down.sql").is_file(),
        "migration sequence advances through 0134 while PRD-091 migration 0083 remains present",
    )
    report.check(
        (ROOT / "infra/migrations/0084_identity_self_service.sql").is_file()
        and (ROOT / "infra/migrations/0084_identity_self_service.down.sql").is_file(),
        "PRD-092 owns reversible migration 0084 with guarded security-fact rollback",
    )
    report.check(
        (ROOT / "infra/migrations/0085_workspace_creation.sql").is_file()
        and (ROOT / "infra/migrations/0085_workspace_creation.down.sql").is_file(),
        "PRD-093 owns reversible migration 0085 with guarded Workspace-fact rollback",
    )
    report.check(
        (ROOT / "infra/migrations/0086_membership_workspace_sessions.sql").is_file()
        and (ROOT / "infra/migrations/0086_membership_workspace_sessions.down.sql").is_file(),
        "PRD-094 owns reversible migration 0086 with guarded authorization-fact rollback",
    )
    report.check(
        (ROOT / "infra/migrations/0087_platform_iam_controlled_support.sql").is_file()
        and (ROOT / "infra/migrations/0087_platform_iam_controlled_support.down.sql").is_file(),
        "PRD-095 owns reversible migration 0087 with guarded platform-security-fact rollback",
    )
    report.check(
        (ROOT / "infra/migrations/0088_platform_workspace_operations.sql").is_file()
        and (ROOT / "infra/migrations/0088_platform_workspace_operations.down.sql").is_file(),
        "PRD-096 owns reversible migration 0088 with guarded operational-fact rollback",
    )
    report.check(
        (ROOT / "infra/migrations/0109_data_residency_retention_legal_hold.sql").is_file()
        and (
            ROOT / "infra/migrations/0109_data_residency_retention_legal_hold.down.sql"
        ).is_file(),
        "PRD-117 owns reversible migration 0109 with guarded governance-fact rollback",
    )
    report.check(
        (ROOT / "infra/migrations/0110_service_account_oauth_machine_identity.sql").is_file()
        and (
            ROOT / "infra/migrations/0110_service_account_oauth_machine_identity.down.sql"
        ).is_file(),
        "PRD-118 owns reversible migration 0110 with guarded machine-identity rollback",
    )
    report.check(
        (
            ROOT
            / "infra/migrations/0111_official_provider_catalog_regional_compliance.sql"
        ).is_file()
        and (
            ROOT
            / "infra/migrations/0111_official_provider_catalog_regional_compliance.down.sql"
        ).is_file(),
        "PRD-119 owns reversible migration 0111 with guarded catalog-fact rollback",
    )
    report.check(
        (ROOT / "infra/migrations/0112_edge_agent_device_identity_mtls.sql").is_file()
        and (
            ROOT / "infra/migrations/0112_edge_agent_device_identity_mtls.down.sql"
        ).is_file(),
        "PRD-120 owns reversible migration 0112 with guarded Edge-identity rollback",
    )
    report.check(
        (
            ROOT
            / "infra/migrations/0113_edge_desired_state_signed_snapshot_lkg.sql"
        ).is_file()
        and (
            ROOT
            / "infra/migrations/0113_edge_desired_state_signed_snapshot_lkg.down.sql"
        ).is_file(),
        "PRD-121 owns reversible migration 0113 with guarded desired-state rollback",
    )
    report.check(
        (
            ROOT
            / "infra/migrations/0114_edge_telemetry_usage_failure_sync.sql"
        ).is_file()
        and (
            ROOT
            / "infra/migrations/0114_edge_telemetry_usage_failure_sync.down.sql"
        ).is_file(),
        "PRD-122 owns reversible migration 0114 with guarded telemetry rollback",
    )
    report.check(
        (ROOT / "infra/migrations/0115_edge_secret_reference_sync.sql").is_file()
        and (
            ROOT / "infra/migrations/0115_edge_secret_reference_sync.down.sql"
        ).is_file(),
        "PRD-123 owns reversible migration 0115 with guarded secret-reference rollback",
    )
    report.check(
        (
            ROOT
            / "infra/migrations/0116_edge_installer_preflight_onboarding.sql"
        ).is_file()
        and (
            ROOT
            / "infra/migrations/0116_edge_installer_preflight_onboarding.down.sql"
        ).is_file(),
        "PRD-124 owns reversible migration 0116 with guarded installer rollback",
    )
    report.check(
        (
            ROOT
            / "infra/migrations/0117_edge_validation_delivery_support_bundles.sql"
        ).is_file()
        and (
            ROOT
            / "infra/migrations/0117_edge_validation_delivery_support_bundles.down.sql"
        ).is_file(),
        "PRD-125 owns reversible migration 0117 with guarded validation rollback",
    )
    report.check(
        (
            ROOT
            / "infra/migrations/0118_edge_fleet_health_controlled_diagnostics.sql"
        ).is_file()
        and (
            ROOT
            / "infra/migrations/0118_edge_fleet_health_controlled_diagnostics.down.sql"
        ).is_file(),
        "PRD-126 owns reversible migration 0118 with guarded Fleet diagnostics rollback",
    )
    report.check(
        (
            ROOT
            / "infra/migrations/0119_edge_fleet_staged_upgrade_rollback.sql"
        ).is_file()
        and (
            ROOT
            / "infra/migrations/0119_edge_fleet_staged_upgrade_rollback.down.sql"
        ).is_file(),
        "PRD-127 owns reversible migration 0119 with guarded Fleet rollout rollback",
    )
    report.check(
        (
            ROOT
            / "infra/migrations/0120_edge_license_capacity_offline_grace.sql"
        ).is_file()
        and (
            ROOT
            / "infra/migrations/0120_edge_license_capacity_offline_grace.down.sql"
        ).is_file(),
        "PRD-128 owns reversible migration 0120 with guarded License rollback",
    )
    report.check(
        (
            ROOT
            / "infra/migrations/0121_edge_backup_restore_dr_drill.sql"
        ).is_file()
        and (
            ROOT
            / "infra/migrations/0121_edge_backup_restore_dr_drill.down.sql"
        ).is_file(),
        "PRD-129 owns reversible migration 0121 with guarded backup rollback",
    )
    report.check(
        (
            ROOT
            / "infra/migrations/0122_edge_offline_bundle_conflict_reconciliation.sql"
        ).is_file()
        and (
            ROOT
            / "infra/migrations/0122_edge_offline_bundle_conflict_reconciliation.down.sql"
        ).is_file(),
        "PRD-130 owns reversible migration 0122 with guarded offline bundle rollback",
    )
    report.check(
        (
            ROOT
            / "infra/migrations/0123_customer_migration_cutover_operations.sql"
        ).is_file()
        and (
            ROOT
            / "infra/migrations/0123_customer_migration_cutover_operations.down.sql"
        ).is_file(),
        "PRD-131 owns reversible migration 0123 with guarded customer migration rollback",
    )
    report.check(
        (
            ROOT
            / "infra/migrations/0124_customer_support_ticket_sla_diagnostics.sql"
        ).is_file()
        and (
            ROOT
            / "infra/migrations/0124_customer_support_ticket_sla_diagnostics.down.sql"
        ).is_file(),
        "PRD-132 owns reversible migration 0124 with guarded customer support rollback",
    )
    report.check(
        (
            ROOT
            / "infra/migrations/0125_public_api_sdk_lifecycle.sql"
        ).is_file()
        and (
            ROOT
            / "infra/migrations/0125_public_api_sdk_lifecycle.down.sql"
        ).is_file(),
        "PRD-133 owns reversible migration 0125 with guarded public contract rollback",
    )
    report.check(
        (ROOT / "infra/migrations/0126_global_account_privacy_dsr.sql").is_file()
        and (ROOT / "infra/migrations/0126_global_account_privacy_dsr.down.sql").is_file(),
        "PRD-134 owns reversible migration 0126 with guarded privacy fact rollback",
    )
    report.check(
        (ROOT / "infra/migrations/0127_compliance_evidence_center.sql").is_file()
        and (ROOT / "infra/migrations/0127_compliance_evidence_center.down.sql").is_file(),
        "PRD-135 owns reversible migration 0127 with guarded compliance fact rollback",
    )
    report.check(
        (ROOT / "infra/migrations/0128_supply_chain_security_provenance.sql").is_file()
        and (ROOT / "infra/migrations/0128_supply_chain_security_provenance.down.sql").is_file(),
        "PRD-136 owns reversible migration 0128 with guarded supply-chain fact rollback",
    )
    report.check(
        (ROOT / "infra/migrations/0129_tenant_isolation_security_assurance.sql").is_file()
        and (
            ROOT / "infra/migrations/0129_tenant_isolation_security_assurance.down.sql"
        ).is_file(),
        "PRD-137 owns reversible migration 0129 with guarded assurance-fact rollback",
    )
    report.check(
        (ROOT / "infra/migrations/0130_saas_release_candidate_launch_gates.sql").is_file()
        and (
            ROOT / "infra/migrations/0130_saas_release_candidate_launch_gates.down.sql"
        ).is_file(),
        "PRD-138 owns reversible migration 0130 with guarded launch-governance rollback",
    )
    report.check(
        (ROOT / "infra/migrations/0131_channel_partner_attribution_management.sql").is_file()
        and (
            ROOT / "infra/migrations/0131_channel_partner_attribution_management.down.sql"
        ).is_file(),
        "PRD-139 owns reversible migration 0131 with guarded channel-fact rollback",
    )
    report.check(
        (ROOT / "infra/migrations/0132_platform_customer_success_renewal.sql").is_file()
        and (
            ROOT / "infra/migrations/0132_platform_customer_success_renewal.down.sql"
        ).is_file(),
        "PRD-140 owns reversible migration 0132 with guarded customer-success rollback",
    )
    report.check(
        (
            ROOT / "infra/migrations/0133_platform_growth_funnel_campaign_attribution.sql"
        ).is_file()
        and (
            ROOT
            / "infra/migrations/0133_platform_growth_funnel_campaign_attribution.down.sql"
        ).is_file(),
        "PRD-141 owns reversible migration 0133 with guarded growth-fact rollback",
    )
    report.check(
        (ROOT / "infra/migrations/0134_platform_channel_deployment_branding.sql").is_file()
        and (
            ROOT / "infra/migrations/0134_platform_channel_deployment_branding.down.sql"
        ).is_file(),
        "PRD-142 owns reversible migration 0134 with guarded branding-fact rollback",
    )

    audit_text = workbook_text + "\n" + _document_texts(REQUIRED_DOCS)
    for name, pattern in SECRET_PATTERNS.items():
        report.check(pattern.search(audit_text) is None, f"no {name} literal in PRD-090 evidence")

    return report


def _is_acyclic(graph: dict[int, set[int]]) -> bool:
    visited: set[int] = set()
    visiting: set[int] = set()

    def visit(node: int) -> bool:
        if node in visiting:
            return False
        if node in visited:
            return True
        visiting.add(node)
        for dependency in graph.get(node, set()):
            if dependency in graph and not visit(dependency):
                return False
        visiting.remove(node)
        visited.add(node)
        return True

    return all(visit(node) for node in graph)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--expected-next")
    args = parser.parse_args()
    report = validate(expected_next=args.expected_next)
    payload = {
        "status": "passed" if report.ok else "failed",
        "checks_passed": len(report.checks),
        "errors": report.errors,
        "next_prd": report.next_prd,
        "roadmap_sha256": hashlib.sha256(NEW_PLAN.read_bytes()).hexdigest() if NEW_PLAN.exists() else None,
    }
    print(json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True))
    return 0 if report.ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
