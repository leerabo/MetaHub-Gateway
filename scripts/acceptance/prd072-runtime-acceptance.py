"""PRD-072 本地真实 HTTP、文件密文与租户权限验收器。

脚本只连接工作区隔离数据库以及 ``127.0.0.1`` 上已经启动的 Gateway、Control 和 Console。
TXT、DOCX、PDF 正文、文件名、登录 Token、密码和 DLP pattern 只在当前调用栈使用；正式证据
只保存迁移版本、HTTP 状态、统一错误码、不可变对象 ID 和计数，不复制任何文件内容或凭据。
"""

from __future__ import annotations

import hashlib
import io
import json
import os
import secrets
import sqlite3
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit
from uuid import uuid4

import httpx
from docx import Document
from metahub_shared.db.security import hash_password

# 验收器只允许访问工作区路径和三个固定回环端口，不能扩展为外部扫描或生产数据操作。
ROOT = Path(__file__).resolve().parents[2]
DEFAULT_DATABASE = ROOT / ".tmp" / "prd072-live" / "metahub-prd072-runtime.db"
DEFAULT_STORAGE = ROOT / ".tmp" / "prd072-live" / "file-storage"
DEFAULT_CONTROL_LOG = ROOT / ".tmp" / "prd072-live" / "control.log"
DEFAULT_GATEWAY_LOG = ROOT / ".tmp" / "prd072-live" / "gateway.log"
DEFAULT_CONSOLE_LOG = ROOT / ".tmp" / "prd072-live" / "console.log"
DEFAULT_ARTIFACT_ROOT = ROOT / ".artifacts" / "prd072-runtime"

CONTROL_URL = "http://127.0.0.1:8001"
GATEWAY_URL = "http://127.0.0.1:8000"
CONSOLE_URL = "http://127.0.0.1:5173"
ADMIN_TENANT_SLUG = "default"
ADMIN_EMAIL = "admin@metahub.local"

# 证据键和递归值扫描共同阻断正文、凭据、文件名、规则与对象定位进入可提交制品。
FORBIDDEN_EVIDENCE_KEYS = frozenset(
    {
        "authorization",
        "content",
        "filename",
        "original_name",
        "password",
        "pattern",
        "provider_key",
        "storage_root",
        "token",
    }
)


class AcceptanceFailure(RuntimeError):
    """只携带固定验收阶段和分类码，禁止异常文本拼入 HTTP 正文或文件内容。"""

    def __init__(self, stage: str, code: str) -> None:
        """构造可公开的失败坐标，不接受自由异常消息。"""

        self.stage = stage
        self.code = code
        super().__init__(f"{stage}:{code}")


def _require(condition: bool, stage: str, code: str) -> None:
    """执行失败关闭断言；失败时只抛出固定阶段和分类。"""

    if not condition:
        # 固定失败坐标既便于 CI 定位，也避免把响应正文、SQL 或文件名带出验收器。
        raise AcceptanceFailure(stage, code)


def _utc_now() -> str:
    """返回数据库和 JSON 均可稳定解析的 UTC 秒级时间。"""

    return datetime.now(UTC).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def _workspace_path(value: str | Path, boundary: Path, *, stage: str) -> Path:
    """解析路径并强制位于指定工作区目录，阻止验收器越过项目边界。"""

    path = Path(value).expanduser().resolve()
    root = boundary.resolve()
    _require(path == root or root in path.parents, stage, "path_outside_workspace")
    return path


def _loopback_url(value: str, *, port: int, stage: str) -> str:
    """校验无凭据、无查询的本机 HTTP 基址，避免验收器访问外部网络。"""

    try:
        parsed = urlsplit(value)
        parsed_port = parsed.port
    # URL 解析异常可能包含调用方文本，统一转换为固定失败坐标。
    except ValueError:
        raise AcceptanceFailure(stage, "url_invalid") from None
    _require(parsed.scheme == "http", stage, "scheme_invalid")
    _require(parsed.hostname in {"127.0.0.1", "localhost"}, stage, "host_invalid")
    _require(parsed_port == port, stage, "port_invalid")
    _require(parsed.username is None and parsed.password is None, stage, "credentials_forbidden")
    _require(not parsed.query and not parsed.fragment, stage, "suffix_invalid")
    _require(parsed.path in {"", "/"}, stage, "path_invalid")
    return value.rstrip("/")


@dataclass(frozen=True, slots=True)
class RuntimeConfig:
    """本地验收配置；网络与所有读写路径都经过固定边界校验。"""

    database: Path
    storage: Path
    control_log: Path
    gateway_log: Path
    console_log: Path
    artifact_root: Path
    control_url: str
    gateway_url: str
    console_url: str
    # 管理员密码只允许从当前进程环境注入，关闭 repr 防止测试失败或调试输出泄露真实值。
    admin_password: str = field(default="", repr=False)

    @classmethod
    def from_environment(cls) -> RuntimeConfig:
        """读取显式环境覆盖，并重新执行工作区路径和回环网络门禁。"""

        # 读取后立即从环境删除，避免后续子进程继承；缺失时失败关闭，绝不回退到仓库内默认密码。
        admin_password = os.environ.pop("METAHUB_PRD072_ADMIN_PASSWORD", "")
        _require(len(admin_password) >= 12, "config.identity", "admin_password_missing")
        return cls(
            database=_workspace_path(
                os.getenv("METAHUB_PRD072_DATABASE", str(DEFAULT_DATABASE)),
                ROOT / ".tmp",
                stage="config.database",
            ),
            storage=_workspace_path(
                os.getenv("METAHUB_PRD072_STORAGE", str(DEFAULT_STORAGE)),
                ROOT / ".tmp",
                stage="config.storage",
            ),
            control_log=_workspace_path(
                os.getenv("METAHUB_PRD072_CONTROL_LOG", str(DEFAULT_CONTROL_LOG)),
                ROOT / ".tmp",
                stage="config.control_log",
            ),
            gateway_log=_workspace_path(
                os.getenv("METAHUB_PRD072_GATEWAY_LOG", str(DEFAULT_GATEWAY_LOG)),
                ROOT / ".tmp",
                stage="config.gateway_log",
            ),
            console_log=_workspace_path(
                os.getenv("METAHUB_PRD072_CONSOLE_LOG", str(DEFAULT_CONSOLE_LOG)),
                ROOT / ".tmp",
                stage="config.console_log",
            ),
            artifact_root=_workspace_path(
                os.getenv("METAHUB_PRD072_ARTIFACT_ROOT", str(DEFAULT_ARTIFACT_ROOT)),
                ROOT / ".artifacts",
                stage="config.artifact_root",
            ),
            control_url=_loopback_url(
                os.getenv("METAHUB_PRD072_CONTROL_URL", CONTROL_URL),
                port=8001,
                stage="config.control_url",
            ),
            gateway_url=_loopback_url(
                os.getenv("METAHUB_PRD072_GATEWAY_URL", GATEWAY_URL),
                port=8000,
                stage="config.gateway_url",
            ),
            console_url=_loopback_url(
                os.getenv("METAHUB_PRD072_CONSOLE_URL", CONSOLE_URL),
                port=5173,
                stage="config.console_url",
            ),
            admin_password=admin_password,
        )


@dataclass(frozen=True, slots=True)
class RunSecrets:
    """当前批次的正文、文件名、规则和临时身份凭据，全部关闭 repr。"""

    suffix: str
    txt_body: bytes = field(repr=False)
    docx_body: bytes = field(repr=False)
    pdf_body: bytes = field(repr=False)
    blocked_body: bytes = field(repr=False)
    dlp_pattern: str = field(repr=False)
    filenames: tuple[str, str, str, str] = field(repr=False)
    security_email: str = field(repr=False)
    other_email: str = field(repr=False)
    # 两个临时账号是真实可登录身份；密码必须逐批随机且关闭 repr，避免中断后留下公开凭据。
    security_password: str = field(repr=False)
    other_password: str = field(repr=False)

    @classmethod
    def create(cls) -> RunSecrets:
        """生成唯一合成内容以及真实可解析 DOCX/PDF，不依赖外部文件或模型服务。"""

        suffix = secrets.token_hex(8)
        txt_text = f"PRD072_RUNTIME_TXT_{suffix} 企业文件治理"
        docx_text = f"PRD072_RUNTIME_DOCX_{suffix} 文本抽取"
        pdf_text = f"PRD072_RUNTIME_PDF_{suffix}"
        pattern = f"PRD072_BLOCK_{secrets.token_hex(10)}"
        return cls(
            suffix=suffix,
            txt_body=txt_text.encode("utf-8"),
            docx_body=_build_docx(docx_text),
            pdf_body=_build_pdf(pdf_text),
            blocked_body=f"runtime blocked {pattern}".encode(),
            dlp_pattern=pattern,
            filenames=("运行验收.txt", "运行验收.docx", "运行验收.pdf", "阻断验收.txt"),
            security_email=f"security-{suffix}@prd072.local",
            other_email=f"admin-{suffix}@prd072-other.local",
            security_password=f"Prd072!{secrets.token_urlsafe(24)}Aa1",
            other_password=f"Prd072!{secrets.token_urlsafe(24)}Bb2",
        )

    def sensitive_values(self, admin_password: str) -> tuple[bytes, ...]:
        """返回必须从日志、密文 Blob 和证据中排除的非空明文值。

        ``admin_password`` 由 RuntimeConfig 的一次性环境读取提供，不能成为模块常量或默认参数。
        """

        values = [
            self.txt_body,
            f"PRD072_RUNTIME_DOCX_{self.suffix}".encode(),
            f"PRD072_RUNTIME_PDF_{self.suffix}".encode("ascii"),
            self.dlp_pattern.encode("ascii"),
            *(name.encode("utf-8") for name in self.filenames),
            self.security_email.encode("utf-8"),
            self.other_email.encode("utf-8"),
            self.security_password.encode("utf-8"),
            self.other_password.encode("utf-8"),
            admin_password.encode("utf-8"),
        ]
        return tuple(value for value in values if value)


@dataclass(frozen=True, slots=True)
class RuntimeIdentities:
    """临时权限和跨租户 HTTP 验收所需的安全资源 ID。"""

    default_tenant_id: str
    security_user_id: str
    other_tenant_id: str
    other_user_id: str
    dlp_rule_id: str


def _build_docx(text: str) -> bytes:
    """使用 python-docx 在内存生成含段落和表格的真实 OOXML 文档。"""

    document = Document()
    document.add_paragraph(text)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "PRD-072 table extraction"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_pdf(text: str) -> bytes:
    """构造带 Helvetica 文本和规范 xref 的最小 PDF，供真实 pypdf 抽取。"""

    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = (
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(stream)).encode("ascii")
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
    )
    payload = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(payload))
        payload.extend(f"{index} 0 obj\n".encode("ascii"))
        payload.extend(body)
        payload.extend(b"\nendobj\n")
    xref = len(payload)
    payload.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    payload.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode(
            "ascii"
        )
    )
    return bytes(payload)


def _prepare_identities(database: Path, secrets_: RunSecrets) -> RuntimeIdentities:
    """在隔离 0064 库创建 SecurityAdmin、第二租户 Admin 和单条 block 规则。"""

    connection = sqlite3.connect(database)
    connection.row_factory = sqlite3.Row
    now = _utc_now()
    security_user_id = f"usr_{uuid4().hex}"
    other_tenant_id = f"tenant_{uuid4().hex}"
    other_user_id = f"usr_{uuid4().hex}"
    dlp_rule_id = f"dlp_{uuid4().hex}"
    try:
        migration = connection.execute(
            "SELECT COUNT(1),MAX(version) FROM schema_migrations WHERE status='applied'"
        ).fetchone()
        _require(tuple(migration) == (64, "0064"), "database", "migration_not_0064")
        owner = connection.execute(
            "SELECT tenant_id FROM users WHERE email=? AND status='active'",
            (ADMIN_EMAIL,),
        ).fetchone()
        _require(owner is not None, "database", "admin_missing")
        default_tenant_id = str(owner["tenant_id"])
        connection.execute(
            """
            INSERT INTO users(
              user_id,tenant_id,email,display_name,password_hash,role,status,created_at,updated_at
            ) VALUES(?,?,?,?,?,'security_admin','active',?,?)
            """,
            (
                security_user_id,
                default_tenant_id,
                secrets_.security_email,
                "PRD072 Security",
                hash_password(secrets_.security_password),
                now,
                now,
            ),
        )
        connection.execute(
            """
            INSERT INTO tenants(tenant_id,slug,name,service_mode,status,created_at,updated_at)
            VALUES(?,?,?,'internal_only','active',?,?)
            """,
            (
                other_tenant_id,
                f"prd072-{secrets_.suffix}",
                "PRD072 Other",
                now,
                now,
            ),
        )
        connection.execute(
            """
            INSERT INTO users(
              user_id,tenant_id,email,display_name,password_hash,role,status,created_at,updated_at
            ) VALUES(?,?,?,?,?,'admin','active',?,?)
            """,
            (
                other_user_id,
                other_tenant_id,
                secrets_.other_email,
                "PRD072 Other Admin",
                hash_password(secrets_.other_password),
                now,
                now,
            ),
        )
        pattern_hash = "sha256:" + hashlib.sha256(
            f"keyword:{secrets_.dlp_pattern}".encode()
        ).hexdigest()
        connection.execute(
            """
            INSERT INTO dlp_rules(
              dlp_rule_id,tenant_id,name,rule_type,pattern,pattern_hash,action,scan_scope,
              status,priority,description,created_by,updated_by,created_at,updated_at
            ) VALUES(?,?,?,'keyword',?,?,'block','input','active',1,'',
                     (SELECT user_id FROM users WHERE email=?),
                     (SELECT user_id FROM users WHERE email=?),?,?)
            """,
            (
                dlp_rule_id,
                default_tenant_id,
                "PRD072 runtime block",
                secrets_.dlp_pattern,
                pattern_hash,
                ADMIN_EMAIL,
                ADMIN_EMAIL,
                now,
                now,
            ),
        )
        connection.commit()
        return RuntimeIdentities(
            default_tenant_id=default_tenant_id,
            security_user_id=security_user_id,
            other_tenant_id=other_tenant_id,
            other_user_id=other_user_id,
            dlp_rule_id=dlp_rule_id,
        )
    # SQL 失败先回滚隔离库，异常正文可能含 pattern 或邮箱，不能原样传播。
    except sqlite3.Error:
        connection.rollback()
        raise AcceptanceFailure("database", "identity_setup_failed") from None
    finally:
        connection.close()


def _login(
    client: httpx.Client,
    *,
    tenant_slug: str,
    email: str,
    password: str,
    stage: str,
) -> str:
    """登录隔离身份并只向当前调用栈返回 Bearer Session Token。"""

    response = client.post(
        "/auth/login",
        json={"tenant_slug": tenant_slug, "email": email, "password": password},
        headers={"X-Request-ID": f"prd072-{stage}-login"},
    )
    _require(response.status_code == 200, stage, "login_failed")
    token = response.json().get("access_token")
    _require(type(token) is str and bool(token), stage, "token_missing")
    # Token 只在当前验收调用栈内传递，证据文件和日志扫描结果均不得持久化其明文。
    return token


def _auth_headers(token: str, request_id: str) -> dict[str, str]:
    """构造验收请求的认证与链路 Header，租户和操作者始终由后端 Session 恢复。"""

    # 仅把临时 Token 交给 loopback HTTP 客户端，返回结构不得进入验收证据或异常文本。
    return {"Authorization": f"Bearer {token}", "X-Request-ID": request_id}


def _assert_private_response(response: httpx.Response, stage: str) -> None:
    """证明文件成功和错误响应都由最外层统一禁止私有缓存。"""

    _require(
        response.headers.get("cache-control") == "no-store, private",
        stage,
        "cache_control_invalid",
    )


def _upload_ready_samples(
    client: httpx.Client,
    token: str,
    secrets_: RunSecrets,
) -> tuple[str, ...]:
    """真实上传 TXT、DOCX、PDF，并复核 ready、抽取器、详情和 request_id。"""

    samples = (
        (secrets_.filenames[0], secrets_.txt_body, "text/plain", "txt-v1"),
        (
            secrets_.filenames[1],
            secrets_.docx_body,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "docx-v1",
        ),
        (secrets_.filenames[2], secrets_.pdf_body, "application/pdf", "pdf-v1"),
    )
    file_ids: list[str] = []
    for index, (name, content, media_type, extractor) in enumerate(samples, start=1):
        request_id = f"prd072-ready-{secrets_.suffix}-{index}"
        response = client.post(
            "/files",
            files={"file": (name, content, media_type)},
            headers=_auth_headers(token, request_id),
            timeout=60.0,
        )
        _assert_private_response(response, f"ready.{index}")
        _require(response.status_code == 201, f"ready.{index}", "status_not_201")
        payload = response.json()
        file_info = payload.get("file", {})
        extract = file_info.get("extract", {})
        _require(payload.get("request_id") == request_id, f"ready.{index}", "request_id_drift")
        _require(file_info.get("status") == "ready", f"ready.{index}", "status_not_ready")
        _require(
            file_info.get("virus_scan_status") == "not_configured",
            f"ready.{index}",
            "virus_status_incorrect",
        )
        _require(extract.get("extractor_version") == extractor, f"ready.{index}", "extractor_drift")
        _require(
            type(extract.get("estimated_tokens")) is int and extract["estimated_tokens"] > 0,
            f"ready.{index}",
            "token_estimate_missing",
        )
        file_id = file_info.get("file_id")
        _require(type(file_id) is str, f"ready.{index}", "file_id_missing")
        detail = client.get(
            f"/files/{file_id}",
            headers=_auth_headers(token, f"prd072-detail-{index}"),
        )
        _assert_private_response(detail, f"detail.{index}")
        _require(detail.status_code == 200, f"detail.{index}", "status_not_200")
        _require(
            detail.json().get("file", {}).get("file_id") == file_id,
            f"detail.{index}",
            "owner_drift",
        )
        file_ids.append(file_id)
    return tuple(file_ids)


def _verify_http_boundaries(
    config: RuntimeConfig,
    secrets_: RunSecrets,
    identities: RuntimeIdentities,
) -> tuple[dict[str, Any], tuple[bytes, ...]]:
    """执行三服务健康、主流程、DLP、权限和跨租户验收，并返回内存 Token canary。"""

    with httpx.Client(base_url=config.control_url, timeout=30.0) as client:
        for stage, url in (
            ("health.gateway", f"{config.gateway_url}/health"),
            ("health.control", f"{config.control_url}/health"),
            ("health.console", config.console_url),
        ):
            response = httpx.get(url, timeout=10.0)
            _require(response.status_code == 200, stage, "not_200")

        admin_token = _login(
            client,
            tenant_slug=ADMIN_TENANT_SLUG,
            email=ADMIN_EMAIL,
            password=config.admin_password,
            stage="admin",
        )
        ready_ids = _upload_ready_samples(client, admin_token, secrets_)
        listed = client.get(
            "/files?limit=25",
            headers=_auth_headers(admin_token, "prd072-list-ready"),
        )
        _assert_private_response(listed, "list.ready")
        _require(listed.status_code == 200, "list.ready", "status_not_200")
        listed_ids = {item.get("file_id") for item in listed.json().get("files", [])}
        _require(set(ready_ids) <= listed_ids, "list.ready", "ready_files_missing")

        blocked_request_id = f"prd072-block-{secrets_.suffix}"
        blocked = client.post(
            "/files",
            files={
                "file": (
                    secrets_.filenames[3],
                    secrets_.blocked_body,
                    "text/plain",
                )
            },
            headers=_auth_headers(admin_token, blocked_request_id),
        )
        _assert_private_response(blocked, "dlp.block")
        _require(blocked.status_code == 403, "dlp.block", "status_not_403")
        _require(blocked.json().get("code") == "file.dlp_blocked", "dlp.block", "code_drift")
        blocked_list = client.get(
            "/files?status=blocked&limit=25",
            headers=_auth_headers(admin_token, "prd072-list-blocked"),
        )
        blocked_items = blocked_list.json().get("files", [])
        blocked_item = next(
            (item for item in blocked_items if item.get("upload_request_id") == blocked_request_id),
            None,
        )
        _require(blocked_item is not None, "dlp.block", "terminal_fact_missing")
        blocked_file_id = blocked_item.get("file_id")
        _require(type(blocked_file_id) is str, "dlp.block", "file_id_missing")

        security_token = _login(
            client,
            tenant_slug=ADMIN_TENANT_SLUG,
            email=secrets_.security_email,
            password=secrets_.security_password,
            stage="security",
        )
        denied = client.post(
            "/files",
            files={"file": ("denied.txt", b"permission body", "text/plain")},
            headers=_auth_headers(security_token, "prd072-security-denied"),
        )
        _assert_private_response(denied, "permission.denied")
        _require(denied.status_code == 403, "permission.denied", "status_not_403")
        _require(
            denied.json().get("code") == "auth.permission_denied",
            "permission.denied",
            "code_drift",
        )

        other_slug = f"prd072-{secrets_.suffix}"
        other_token = _login(
            client,
            tenant_slug=other_slug,
            email=secrets_.other_email,
            password=secrets_.other_password,
            stage="other",
        )
        other_list = client.get(
            "/files",
            headers=_auth_headers(other_token, "prd072-other-list"),
        )
        hidden = client.get(
            f"/files/{ready_ids[0]}",
            headers=_auth_headers(other_token, "prd072-other-detail"),
        )
        _require(other_list.status_code == 200, "tenant.list", "status_not_200")
        _require(other_list.json().get("files") == [], "tenant.list", "cross_tenant_rows")
        _require(hidden.status_code == 404, "tenant.detail", "status_not_404")
        _require(hidden.json().get("code") == "file.not_found", "tenant.detail", "code_drift")

    # Token canary 只交给本进程日志扫描器，正式 HTTP 证据仅保留状态、对象 ID 和固定错误码。
    return (
        {
            "service_http_status": {"gateway": 200, "control": 200, "console": 200},
            "ready_file_ids": list(ready_ids),
            "blocked_file_id": blocked_file_id,
            "error_codes": {
                "dlp": "file.dlp_blocked",
                "permission": "auth.permission_denied",
                "cross_tenant": "file.not_found",
            },
        },
        tuple(token.encode("utf-8") for token in (admin_token, security_token, other_token)),
    )


def _verify_database_and_storage(
    config: RuntimeConfig,
    identities: RuntimeIdentities,
    http_evidence: dict[str, Any],
    secrets_: RunSecrets,
) -> dict[str, Any]:
    """复核 0064 owner、事件、零阻断 Blob、密文对象和数据库完整性。"""

    ready_ids = tuple(http_evidence["ready_file_ids"])
    blocked_id = str(http_evidence["blocked_file_id"])
    connection = sqlite3.connect(config.database)
    connection.row_factory = sqlite3.Row
    try:
        ready_placeholders = ",".join("?" for _ in ready_ids)
        ready_rows = connection.execute(
            "SELECT * FROM uploaded_files "
            f"WHERE tenant_id=? AND file_id IN ({ready_placeholders})",
            (identities.default_tenant_id, *ready_ids),
        ).fetchall()
        blocked_row = connection.execute(
            "SELECT * FROM uploaded_files WHERE tenant_id=? AND file_id=?",
            (identities.default_tenant_id, blocked_id),
        ).fetchone()
        _require(len(ready_rows) == 3, "database.ready", "row_count_invalid")
        _require(
            all(row["status"] == "ready" and row["original_object_key"] for row in ready_rows),
            "database.ready",
            "owner_state_invalid",
        )
        _require(blocked_row is not None, "database.blocked", "owner_missing")
        _require(
            blocked_row["status"] == "blocked" and blocked_row["original_object_key"] is None,
            "database.blocked",
            "blob_owner_present",
        )
        file_ids = (*ready_ids, blocked_id)
        file_placeholders = ",".join("?" for _ in file_ids)
        extract_count = connection.execute(
            "SELECT COUNT(1) FROM file_extracts "
            f"WHERE tenant_id=? AND file_id IN ({file_placeholders})",
            (identities.default_tenant_id, *file_ids),
        ).fetchone()[0]
        event_count = connection.execute(
            "SELECT COUNT(1) FROM file_events "
            f"WHERE tenant_id=? AND file_id IN ({file_placeholders})",
            (identities.default_tenant_id, *file_ids),
        ).fetchone()[0]
        match_count = connection.execute(
            "SELECT COUNT(1) FROM file_dlp_matches WHERE tenant_id=? AND file_id=?",
            (identities.default_tenant_id, blocked_id),
        ).fetchone()[0]
        _require(
            (extract_count, event_count, match_count) == (4, 8, 1),
            "database",
            "fact_counts_invalid",
        )
        integrity = connection.execute("PRAGMA integrity_check").fetchone()[0]
        foreign_keys = len(connection.execute("PRAGMA foreign_key_check").fetchall())
        _require(integrity == "ok" and foreign_keys == 0, "database", "integrity_invalid")
    # 验收查询异常可能包含数据库路径或 SQL，统一转换为固定失败坐标。
    except sqlite3.Error:
        raise AcceptanceFailure("database", "verification_failed") from None
    finally:
        connection.close()

    blobs = sorted(config.storage.rglob("*.bin"))
    _require(len(blobs) == 6, "storage", "blob_count_invalid")
    # 阻断 pattern 也属于日志禁用值；ready 正文与全部文件名必须从数据库和 Blob 明文中消失。
    storage_bytes = [path.read_bytes() for path in blobs]
    for value in secrets_.sensitive_values(config.admin_password):
        _require(all(value not in blob for blob in storage_bytes), "storage", "plaintext_found")
    database_bytes = config.database.read_bytes()
    for value in (
        secrets_.txt_body,
        f"PRD072_RUNTIME_DOCX_{secrets_.suffix}".encode(),
        f"PRD072_RUNTIME_PDF_{secrets_.suffix}".encode("ascii"),
        *(name.encode("utf-8") for name in secrets_.filenames),
    ):
        _require(value not in database_bytes, "database", "file_plaintext_found")
    return {
        "migration_count": 64,
        "migration_max": "0064",
        "integrity_check": "ok",
        "foreign_key_violations": 0,
        "uploaded_file_count": 4,
        "file_extract_count": 4,
        "file_event_count": 8,
        "file_dlp_match_count": 1,
        "encrypted_blob_count": 6,
    }


def _cleanup_runtime_identities(database: Path, identities: RuntimeIdentities) -> dict[str, int]:
    """禁用临时规则/用户并撤销其 Session，且回读数据库证明真实清理结果。"""

    connection = sqlite3.connect(database)
    now = _utc_now()
    try:
        rule_cursor = connection.execute(
            "UPDATE dlp_rules SET status='disabled',updated_at=? "
            "WHERE dlp_rule_id=? AND status='active'",
            (now, identities.dlp_rule_id),
        )
        user_cursor = connection.execute(
            "UPDATE users SET status='disabled',updated_at=? "
            "WHERE user_id IN (?,?) AND status='active'",
            (now, identities.security_user_id, identities.other_user_id),
        )
        session_cursor = connection.execute(
            """
            UPDATE sessions SET status='revoked',revoked_at=?,updated_at=?
            WHERE user_id IN (?,?) AND status='active'
            """,
            (now, now, identities.security_user_id, identities.other_user_id),
        )
        disabled_rule_count = connection.execute(
            "SELECT COUNT(1) FROM dlp_rules WHERE dlp_rule_id=? AND status='disabled'",
            (identities.dlp_rule_id,),
        ).fetchone()[0]
        disabled_user_count = connection.execute(
            "SELECT COUNT(1) FROM users WHERE user_id IN (?,?) AND status='disabled'",
            (identities.security_user_id, identities.other_user_id),
        ).fetchone()[0]
        active_session_count = connection.execute(
            "SELECT COUNT(1) FROM sessions WHERE user_id IN (?,?) AND status='active'",
            (identities.security_user_id, identities.other_user_id),
        ).fetchone()[0]
        # 固定计数会把错误 ID 或清理失效伪装成成功；必须同时校验 UPDATE 影响行和持久终态。
        _require(rule_cursor.rowcount == 1, "cleanup", "rule_update_count_invalid")
        _require(user_cursor.rowcount == 2, "cleanup", "user_update_count_invalid")
        _require(disabled_rule_count == 1, "cleanup", "rule_terminal_state_invalid")
        _require(disabled_user_count == 2, "cleanup", "user_terminal_state_invalid")
        _require(active_session_count == 0, "cleanup", "session_terminal_state_invalid")
        connection.commit()
        # 清理证据只保留计数，不返回临时用户、规则或 Session 的可识别字段。
        return {
            "disabled_rules": rule_cursor.rowcount,
            "disabled_users": user_cursor.rowcount,
            "revoked_sessions": session_cursor.rowcount,
        }
    # 业务终态检查失败同样必须回滚，不能留下只完成一半的身份清理事务。
    except AcceptanceFailure:
        connection.rollback()
        raise
    # 清理失败不能伪装成验收成功，也不能暴露 SQL 或身份明文。
    except sqlite3.Error:
        connection.rollback()
        raise AcceptanceFailure("cleanup", "database_cleanup_failed") from None
    finally:
        connection.close()


def _sensitive_log_encodings(value: bytes) -> tuple[bytes, ...]:
    """生成日志重定向可能采用的敏感值字节表示。

    参数：验收器内部统一保存的 UTF-8 敏感值。返回：去重后的 UTF-8、UTF-16LE 与
    UTF-16BE 字节序列。Windows PowerShell 的 ``*>`` 默认可能把子进程输出转成
    UTF-16；只扫描 UTF-8 会产生假阴性，因此三种表示必须等价失败关闭。
    """

    try:
        text = value.decode("utf-8", errors="strict")
    # 当前 canary 都由 UTF-8 文本产生；若未来调用方传入二进制值，至少保留原字节扫描，
    # 不能因辅助解码失败而跳过这一项或让整个验收器意外崩溃。
    except UnicodeDecodeError:
        return (value,)
    candidates = (value, text.encode("utf-16-le"), text.encode("utf-16-be"))
    return tuple(dict.fromkeys(candidates))


def _scan_logs(config: RuntimeConfig, sensitive_values: tuple[bytes, ...]) -> dict[str, int]:
    """扫描三服务日志，证明文件名、正文、DLP pattern、密码和 Session Token 命中为零。"""

    scanned = 0
    for path in (config.gateway_log, config.control_log, config.console_log):
        if not path.exists():
            continue
        payload = path.read_bytes()
        scanned += 1
        for value in sensitive_values:
            for encoded_value in _sensitive_log_encodings(value):
                _require(encoded_value not in payload, "logs", "sensitive_value_found")
    _require(scanned == 3, "logs", "expected_logs_missing")
    return {"files_scanned": scanned, "sensitive_matches": 0}


def _assert_safe_evidence(value: object, sensitive_values: tuple[bytes, ...]) -> None:
    """递归拒绝禁用键和任意敏感明文进入正式 JSON 证据。"""

    if isinstance(value, dict):
        for key, item in value.items():
            _require(str(key).lower() not in FORBIDDEN_EVIDENCE_KEYS, "evidence", "forbidden_key")
            _assert_safe_evidence(item, sensitive_values)
        return
    if isinstance(value, list):
        for item in value:
            _assert_safe_evidence(item, sensitive_values)
        return
    if isinstance(value, str):
        encoded = value.encode("utf-8")
        _require(
            all(secret not in encoded for secret in sensitive_values),
            "evidence",
            "secret_value",
        )


def _write_evidence(config: RuntimeConfig, payload: dict[str, Any]) -> Path:
    """校验后以同目录临时文件原子发布脱敏 Runtime 证据。"""

    config.artifact_root.mkdir(parents=True, exist_ok=True)
    run_directory = config.artifact_root / str(payload["run_id"])
    run_directory.mkdir(parents=False, exist_ok=False)
    destination = run_directory / "runtime-acceptance.json"
    temporary = run_directory / ".runtime-acceptance.json.tmp"
    encoded = json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True).encode("utf-8")
    with temporary.open("xb") as handle:
        handle.write(encoded)
        handle.flush()
        os.fsync(handle.fileno())
    os.replace(temporary, destination)
    return destination


def run() -> Path:
    """执行完整本地验收并返回唯一脱敏证据路径。"""

    config = RuntimeConfig.from_environment()
    secrets_ = RunSecrets.create()
    _require(config.database.is_file(), "database", "file_missing")
    _require(config.storage.is_dir(), "storage", "directory_missing")
    identities = _prepare_identities(config.database, secrets_)
    try:
        http_evidence, session_token_canaries = _verify_http_boundaries(
            config,
            secrets_,
            identities,
        )
        database_evidence = _verify_database_and_storage(
            config,
            identities,
            http_evidence,
            secrets_,
        )
    finally:
        # 身份一旦提交，后续 HTTP、数据库、日志或证据阶段的任何异常都必须先收敛临时权限。
        cleanup = _cleanup_runtime_identities(config.database, identities)
    sensitive_values = secrets_.sensitive_values(config.admin_password) + session_token_canaries
    log_evidence = _scan_logs(config, sensitive_values)
    run_id = f"prd072_{datetime.now(UTC).strftime('%Y%m%dT%H%M%SZ')}_{secrets_.suffix}"
    payload: dict[str, Any] = {
        "schema": "metahub.prd072.runtime-acceptance.v1",
        "status": "passed",
        "run_id": run_id,
        "completed_at": _utc_now(),
        "http": http_evidence,
        "database": database_evidence,
        "logs": log_evidence,
        "cleanup": cleanup,
        "all_108_prds_complete": False,
    }
    _assert_safe_evidence(payload, sensitive_values)
    destination = _write_evidence(config, payload)
    _require(
        all(value not in destination.read_bytes() for value in sensitive_values),
        "evidence",
        "plaintext_found_after_write",
    )
    return destination


def main() -> int:
    """运行验收并只打印状态与工作区内证据相对路径。"""

    try:
        destination = run()
    # 固定失败只输出阶段和分类，不打印异常链、HTTP 正文或文件内容。
    except AcceptanceFailure as exc:
        print(json.dumps({"status": "failed", "stage": exc.stage, "code": exc.code}))
        return 1
    print(
        json.dumps(
            {
                "status": "passed",
                "artifact": destination.relative_to(ROOT).as_posix(),
            }
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
